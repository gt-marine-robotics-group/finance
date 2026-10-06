"""Bill requests must stop before Engage when screenshot proof is missing."""

from pathlib import Path
import hashlib
from unittest.mock import MagicMock

import openpyxl
import pandas as pd
import pytest

from mrg_finance import automation, automation_screenshots
from mrg_finance.screenshot_capture import (
    BrowserChallenge, capture_evidence, evidence_metadata_path, validate_evidence,
)


URL = "https://www.digikey.com/en/products/detail/raspberry-pi/SC0194-9/10258781"


def browser():
    driver = MagicMock()
    driver.current_url = URL
    driver.title = "SC0194(9) Raspberry Pi"
    driver.find_element.return_value.text = "Raspberry Pi $100.00 Add to Cart"
    driver.find_elements.return_value = []
    driver.execute_script.return_value = "complete"
    driver.execute_cdp_cmd.side_effect = RuntimeError("No CDP")
    driver.save_screenshot.side_effect = lambda p: Path(p).write_bytes(b"product image") or True
    return driver


def test_bill_reuses_only_checked_source_and_unchanged_image(monkeypatch, tmp_path):
    monkeypatch.setattr(automation, "SCREENSHOT_DIR", str(tmp_path))
    path = tmp_path / "Test Bill" / "Part.png"
    capture_evidence(browser(), path, source_url=URL)
    automation.require_bill_evidence([("Part", URL)], "Test Bill")
    assert automation._verified_screenshot("Part", "Test Bill", URL) == str(path)
    assert not validate_evidence(path, source_url="https://www.amazon.com/dp/B012345678")
    path.write_bytes(b"replaced with challenge image")
    with pytest.raises(ValueError, match="before Engage"):
        automation.require_bill_evidence([("Part", URL)], "Test Bill")


def test_challenge_during_capture_invalidates_image_and_previous_proof(tmp_path):
    path = tmp_path / "Part.png"
    driver = browser()
    capture_evidence(driver, path)
    assert validate_evidence(path, source_url=URL)
    def screenshot(p):
        Path(p).write_bytes(b"challenge image")
        driver.title = "Robot check"
        driver.find_element.return_value.text = "Verify you are human"
        return True
    driver.save_screenshot.side_effect = screenshot
    with pytest.raises(BrowserChallenge):
        capture_evidence(driver, path)
    assert not path.exists()
    assert not validate_evidence(path, source_url=URL)
    assert evidence_metadata_path(path).exists()
    assert list((tmp_path / "challenges").glob("*.png"))
    path.write_bytes(b"product image")  # older image restored by a cloud download
    assert not validate_evidence(path, source_url=URL)


def test_noninteractive_verification_can_finish_before_screenshot(tmp_path):
    driver = browser()
    driver.title = "DigiKey"
    # Initial challenge, then the vendor's own verification finishes naturally.
    driver.find_element.side_effect = [
        MagicMock(text="Performing security verification"),
        MagicMock(text="Raspberry Pi $100.00"),
        MagicMock(text="Raspberry Pi $100.00"),
    ]
    path = tmp_path / "Part.png"
    capture_evidence(driver, path, settle_timeout=1)
    assert validate_evidence(path, source_url=URL)
    assert not (tmp_path / "challenges").exists()


@pytest.mark.parametrize("capture", [False, True])
def test_bill_main_stops_before_engage_if_capture_declined_or_blocked(monkeypatch, tmp_path, capture):
    path = tmp_path / "bill.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bills"
    ws.append(["Bill Title", "Bill No.", "Item Name", "Link", "Cost", "Quantity"])
    ws.append(["Test Bill", "344042", "Part", URL, 63.35, 5])
    wb.save(path)
    wb.close()
    monkeypatch.setattr(automation, "SCREENSHOT_DIR", str(tmp_path / "screenshots"))
    monkeypatch.setattr(automation, "USERNAME", "test-only")
    monkeypatch.setattr(automation, "PASSWORD", "test-only")
    monkeypatch.setattr(automation, "BILL_NO", "")
    monkeypatch.setattr(automation, "BILL_URL", "")
    monkeypatch.setattr("sys.argv", ["bill-request", "--bill", "Test Bill", "--excel-path", str(path)])
    monkeypatch.setattr("builtins.input", lambda p: "y" if capture else "n")
    monkeypatch.setattr(automation.time, "sleep", lambda *a: None)
    chrome = MagicMock(return_value=browser())
    monkeypatch.setattr(automation.webdriver, "Chrome", chrome)
    monkeypatch.setattr("selenium.webdriver.chrome.service.Service", MagicMock())
    from mrg_finance import screenshot_capture
    monkeypatch.setattr(screenshot_capture, "capture_evidence", MagicMock(side_effect=BrowserChallenge("Unresolved challenge")))
    # automation.py's direct-module imports are used by the CLI subprocess.
    import price_scraper, automation_screenshots as direct_screenshots
    monkeypatch.setattr(price_scraper, "dismiss_popups_and_interstitials", lambda d: None)
    monkeypatch.setattr(direct_screenshots, "sync_screenshots_to_sharepoint", lambda folder: None)
    with pytest.raises(ValueError, match="before Engage"):
        automation.main()
    assert chrome.call_count == int(capture)  # capture browser only; no Engage browser
    if capture:
        chrome.return_value.quit.assert_called_once()
        assert any(a.startswith("--user-data-dir=") for a in chrome.call_args.kwargs["options"].arguments)


def test_legacy_image_without_proof_is_not_bill_evidence(monkeypatch, tmp_path):
    monkeypatch.setattr(automation, "SCREENSHOT_DIR", str(tmp_path))
    folder = tmp_path / "Test Bill"
    folder.mkdir()
    (folder / "Part.png").write_bytes(b"unchecked older image")
    assert automation._find_screenshot("Part", "Test Bill")
    assert automation._verified_screenshot("Part", "Test Bill", URL) is None
    with pytest.raises(ValueError, match="Part"):
        automation.require_bill_evidence([("Part", URL)], "Test Bill")


@pytest.mark.parametrize("numbers", [[""], ["344042", ""], ["344042", "344043"],
                                    ["https://example.com/edit/344042"], [0], ["344042.5"]])
def test_bill_missing_or_conflicting_draft_ids_stop_before_credentials_or_chrome(monkeypatch, tmp_path, numbers):
    path = tmp_path / "bill.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bills"
    ws.append(["Bill Title", "Bill No.", "Item Name", "Link"])
    for index, number in enumerate(numbers):
        ws.append(["Test Bill", number, f"Part {index}", URL])
    wb.save(path)
    wb.close()
    monkeypatch.setattr("sys.argv", ["bill-request", "--excel-path", str(path), "--bill", "Test Bill"])
    monkeypatch.setattr(automation, "USERNAME", "")
    monkeypatch.setattr(automation, "PASSWORD", "")
    monkeypatch.setattr(automation, "BILL_URL", "https://example.com/stale-draft")
    unexpected = MagicMock(side_effect=AssertionError("Must validate the workbook reference first"))
    monkeypatch.setattr("builtins.input", unexpected)
    monkeypatch.setattr(automation.getpass, "getpass", unexpected)
    chrome = MagicMock()
    monkeypatch.setattr(automation.webdriver, "Chrome", chrome)
    with pytest.raises(automation.BillDraftRequired, match="Create and save the bill draft in Engage first"):
        automation.main()
    unexpected.assert_not_called()
    chrome.assert_not_called()


def test_bill_draft_check_accepts_excel_numeric_ids_and_ignores_separator_rows():
    rows = pd.DataFrame([
        {"Item": "Part A", "Bill Number": 344042.0},
        {"Item": "Part B", "Bill Number": "344042"},
        {"Item": "", "Bill Number": ""},
    ])
    assert automation.require_engage_draft_number(rows, "Test Bill") == "344042"


@pytest.fixture
def screenshot_workbook(tmp_path, monkeypatch):
    path = tmp_path / "bill.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bills"
    ws.append(["Bill Name", "Item", "Product URL", "Unit Cost"])
    ws.append(["Request 1", "", "", ""])
    ws.append(["Liquid - RobotX", "Non-bill Item - Update Manually", "", ""])
    ws.append(["Misc", "Misc", "", ""])
    ws.append(["Test Bill", "Part", URL, 63.35])
    ws.append(["Test Bill", "", "", ""])
    ws.append(["Test Bill", "Missing link", "", 20])
    ws.append(["Test Bill", "Malformed link", "http://[", 30])
    wb.save(path)
    wb.close()
    monkeypatch.setattr(automation_screenshots, "SAVE_FOLDER", str(tmp_path / "shots"))
    monkeypatch.setattr(automation_screenshots, "Service", MagicMock())
    monkeypatch.setattr(automation_screenshots, "sync_screenshots_to_sharepoint", MagicMock())
    return path


def test_screenshot_main_filters_menu_retries_selection_and_audits_actual_items(screenshot_workbook, monkeypatch, tmp_path, capsys):
    path = screenshot_workbook
    before = hashlib.sha256(path.read_bytes()).digest()
    driver = browser()
    chrome = MagicMock(return_value=driver)
    monkeypatch.setattr(automation_screenshots.webdriver, "Chrome", chrome)
    monkeypatch.setattr(automation_screenshots, "dismiss_popups", lambda d: None)
    monkeypatch.setattr(automation_screenshots, "extract_price_from_page", lambda *a: ("$100.00", "high"))
    from mrg_finance import screenshot_capture
    navigate = MagicMock()
    monkeypatch.setattr(screenshot_capture, "navigate_for_evidence", navigate)
    monkeypatch.setattr("sys.argv", ["screenshots", "--excel-path", str(path), "--headless"])
    answers = iter(["", "0", "99", "test bill"])
    monkeypatch.setattr("builtins.input", lambda p: next(answers))
    assert automation_screenshots.main() == 1  # Two real product rows need URLs.
    output = capsys.readouterr().out
    assert "  1. Test Bill" in output
    assert not any(t in output for t in ("Request 1", "Liquid - RobotX", "Misc"))
    assert output.count("Choose a listed bill") == 3
    assert "Screenshot mode: headless Chrome" in output
    assert "Items: 3 | Need attention: 2" in output
    folder = tmp_path / "shots" / "Test Bill"
    automation_screenshots.sync_screenshots_to_sharepoint.assert_called_once_with(folder)
    audit = pd.read_csv(folder / "screenshot_audit.csv")
    assert list(audit["Status"]) == ["captured", "missing_url", "invalid_url"]
    assert audit.iloc[0]["Quoted Unit Cost"] == 100
    assert validate_evidence(folder / "Part.png", source_url=URL)
    navigate.assert_called_once_with(driver, URL)
    driver.quit.assert_called_once()
    assert "--headless=new" in chrome.call_args.kwargs["options"].arguments
    assert hashlib.sha256(path.read_bytes()).digest() == before


def test_screenshot_upload_scopes_source_and_destination_to_current_bill(monkeypatch, tmp_path):
    folder = tmp_path / "shots" / "Test Bill"
    folder.mkdir(parents=True)
    image = folder / "Part.png"
    image.write_bytes(b"product evidence")
    unrelated = folder.parent / "Other Order"
    unrelated.mkdir()
    (unrelated / "Budget_vs_Quoted.xlsx").write_bytes(b"unrelated report")
    run = MagicMock(return_value=MagicMock(returncode=0))
    monkeypatch.setattr("subprocess.run", run)
    automation_screenshots.sync_screenshots_to_sharepoint(folder)
    args = run.call_args.args[0]
    assert args[-2:] == [str(folder), "onedrive:OPS-1 Operations/FY27 Finances/screenshots/Test Bill"]
    assert image.read_bytes() == b"product evidence"


def test_screenshot_upload_failure_is_concise_and_preserves_evidence(monkeypatch, tmp_path, capsys):
    image = tmp_path / "Part.png"
    image.write_bytes(b"product evidence")
    error = 'ERROR : Failed to copy: HTTP error 404: The upload session was not found'
    monkeypatch.setattr("subprocess.run", MagicMock(return_value=MagicMock(returncode=1, stderr="\n".join([error] * 12))))
    automation_screenshots.sync_screenshots_to_sharepoint(tmp_path)
    output = capsys.readouterr().out
    assert output.count(error) == 1
    assert "Local evidence remains saved" in output
    assert image.read_bytes() == b"product evidence"


def test_bill_budget_timeout_recovers_in_same_window_before_any_items_change(monkeypatch, tmp_path):
    from selenium.common.exceptions import TimeoutException
    from mrg_finance import engage_bill_lookup, engage_fields
    url = engage_bill_lookup.build_bill_url("344042")
    driver = MagicMock(current_url=url)
    navigate = MagicMock(side_effect=[TimeoutException(), "Budget Section: B03"])
    monkeypatch.setattr(engage_bill_lookup, "open_budget_view", navigate)
    diagnostic = MagicMock()
    monkeypatch.setattr(engage_fields, "save_engage_diagnostic", diagnostic)

    def retry(prompt):
        driver.get.assert_not_called()
        driver.quit.assert_not_called()
        return ""

    monkeypatch.setattr("builtins.input", retry)
    assert automation.prepare_bill_budget(driver, url, tmp_path) == "Budget Section: B03"
    assert navigate.call_count == 2
    assert all(call.kwargs == {"require_editable": True} for call in navigate.call_args_list)
    diagnostic.assert_called_once_with(driver, tmp_path)
    driver.quit.assert_not_called()


def test_bill_budget_wrong_request_can_be_cancelled_without_mutation(monkeypatch, tmp_path):
    from mrg_finance import engage_bill_lookup, engage_fields
    driver = MagicMock(current_url=engage_bill_lookup.build_bill_url("999"))
    navigate = MagicMock()
    monkeypatch.setattr(engage_bill_lookup, "open_budget_view", navigate)
    monkeypatch.setattr(engage_fields, "save_engage_diagnostic", MagicMock())
    monkeypatch.setattr("builtins.input", lambda prompt: "cancel")
    with pytest.raises(SystemExit, match="Bill preparation cancelled"):
        automation.prepare_bill_budget(driver, engage_bill_lookup.build_bill_url("344042"), tmp_path)
    navigate.assert_not_called()
    driver.get.assert_not_called()
    driver.quit.assert_not_called()


def test_screenshot_navigation_failure_invalidates_old_proof_and_saves_audit(screenshot_workbook, monkeypatch, tmp_path):
    from mrg_finance import screenshot_capture
    from selenium.common.exceptions import TimeoutException
    driver = browser()
    shot = tmp_path / "shots" / "Test Bill" / "Part.png"
    capture_evidence(driver, shot, source_url=URL)
    assert validate_evidence(shot, source_url=URL)
    monkeypatch.setattr(automation_screenshots.webdriver, "Chrome", lambda **kw: driver)
    navigate = MagicMock(side_effect=TimeoutException("Page did not load"))
    monkeypatch.setattr(screenshot_capture, "navigate_for_evidence", navigate)
    monkeypatch.setattr("sys.argv", ["screenshots", "--excel-path", str(screenshot_workbook),
                                    "--bill", "Test Bill", "--interactive"])
    assert automation_screenshots.main() == 1
    assert navigate.call_count == 3
    assert not validate_evidence(shot, source_url=URL)
    audit = pd.read_csv(shot.parent / "screenshot_audit.csv")
    assert audit.iloc[0]["Status"] == "failed"
    assert "Page did not load" in audit.iloc[0]["Error"]
    driver.quit.assert_called_once()


def test_screenshot_menu_cancel_does_not_start_browser(screenshot_workbook, monkeypatch):
    chrome = MagicMock()
    monkeypatch.setattr(automation_screenshots.webdriver, "Chrome", chrome)
    monkeypatch.setattr("sys.argv", ["screenshots", "--excel-path", str(screenshot_workbook)])
    monkeypatch.setattr("builtins.input", lambda p: "cancel")
    assert automation_screenshots.main() == 0
    chrome.assert_not_called()
