"""DigiKey challenge recovery, quantity handling, and durable cart references."""

import json
from pathlib import Path
from unittest.mock import MagicMock

import openpyxl
import pytest
from selenium.common.exceptions import ElementClickInterceptedException, TimeoutException
from selenium.webdriver.common.by import By

from mrg_finance import digikey_cart, purchase_cart, share_a_cart
from mrg_finance.screenshot_capture import BrowserChallenge, capture_evidence
from mrg_finance.spreadsheet_utils import update_order_table_links
from test_purchase_replacements import workbook


ITEM = {"item_name": "Raspberry Pi", "quantity": 5, "cost": 63.35,
        "link": "https://www.digikey.com/en/products/detail/raspberry-pi/SC0194-9/10258781"}
URL = "https://share-a-cart.com/get/DIGI01"


class ImmediateWait:
    def __init__(self, *a, **kw):
        pass

    def until(self, condition):
        result = condition(self.driver)
        if not result:
            raise TimeoutException()
        return result


def setup_product(monkeypatch):
    driver, quantity, button = MagicMock(), MagicMock(), MagicMock()
    attrs = {"value": "1", "max": "150"}
    quantity.get_attribute.side_effect = attrs.get
    quantity.send_keys.side_effect = lambda value: attrs.update(value=value) if value == "5" else None
    driver.find_elements.side_effect = lambda by, selector: [button] if by == By.XPATH else [quantity] if selector.startswith("input[") else []
    wait = ImmediateWait()
    wait.driver = driver
    monkeypatch.setattr(digikey_cart, "WebDriverWait", lambda *a, **kw: wait)
    monkeypatch.setattr(digikey_cart, "navigate_for_evidence", MagicMock())
    monkeypatch.setattr(digikey_cart.price_scraper, "dismiss_popups_and_interstitials", MagicMock())
    monkeypatch.setattr(digikey_cart, "capture_evidence", MagicMock())
    monkeypatch.setattr(digikey_cart, "save_page_screenshot", lambda d, p, **kw: str(p))
    return driver, quantity, button, attrs


def test_digikey_adds_requested_quantity_once(monkeypatch, tmp_path):
    driver, quantity, button, attrs = setup_product(monkeypatch)
    counts = iter([0, 5])
    monkeypatch.setattr(digikey_cart, "_cart_count", lambda d: next(counts))
    digikey_cart.add_digikey_item(driver, ITEM, tmp_path)
    assert attrs["value"] == "5"
    button.click.assert_called_once()
    driver.quit.assert_not_called()


@pytest.mark.parametrize("reason", ["blocked", "no_confirmation", "quantity_reset"])
def test_unconfirmed_digikey_add_keeps_browser_for_manual_recovery(monkeypatch, tmp_path, reason):
    driver, quantity, button, attrs = setup_product(monkeypatch)
    monkeypatch.setattr(digikey_cart, "_cart_count", lambda d: 0)
    if reason == "blocked":
        button.click.side_effect = ElementClickInterceptedException()
    if reason == "quantity_reset":
        quantity.send_keys.side_effect = None
    def finish(prompt):
        driver.quit.assert_not_called()
        assert "Finish adding" in prompt
        assert button.click.call_count <= 1
        return ""
    monkeypatch.setattr("builtins.input", finish)
    digikey_cart.add_digikey_item(driver, ITEM, tmp_path)
    assert button.click.call_count == (0 if reason == "quantity_reset" else 1)


def test_digikey_quantity_limit_stops_without_clicking(monkeypatch, tmp_path):
    driver, quantity, button, attrs = setup_product(monkeypatch)
    attrs["max"] = "3"
    with pytest.raises(ValueError, match="requested 5"):
        digikey_cart.add_digikey_item(driver, ITEM, tmp_path)
    button.click.assert_not_called()


def test_digikey_uses_main_buy_box_quantity_among_related_product_fields(monkeypatch, tmp_path):
    driver, quantity, button, attrs = setup_product(monkeypatch)
    scope, related = MagicMock(), MagicMock()
    button.find_elements.return_value = [scope]
    scope.find_elements.return_value = [quantity]
    driver.find_elements.side_effect = lambda by, selector: [button] if by == By.XPATH else [quantity, related] if selector.startswith("input[") else []
    counts = iter([0, 5])
    monkeypatch.setattr(digikey_cart, "_cart_count", lambda d: next(counts))
    digikey_cart.add_digikey_item(driver, ITEM, tmp_path)
    assert attrs["value"] == "5"
    button.click.assert_called_once()
    related.clear.assert_not_called()
    related.send_keys.assert_not_called()


@pytest.mark.parametrize("empty", [False, True])
def test_existing_digikey_cart_is_not_duplicated(monkeypatch, tmp_path, empty):
    driver = MagicMock()
    driver.find_element.return_value.text = "Your shopping cart is empty." if empty else "Your Cart: Raspberry Pi x5"
    monkeypatch.setattr(digikey_cart, "navigate_for_evidence", MagicMock())
    monkeypatch.setattr(digikey_cart, "capture_evidence", MagicMock())
    add = MagicMock()
    monkeypatch.setattr(digikey_cart, "add_digikey_item", add)
    digikey_cart.prepare_digikey_items(driver, [ITEM], tmp_path)
    assert add.call_count == int(empty)


@pytest.mark.parametrize("body", ["Performing security verification", ""])
def test_cloudflare_challenge_is_saved_and_requires_manual_completion(tmp_path, body):
    driver = MagicMock()
    driver.current_url = ITEM["link"]
    driver.title = "Just a moment..."
    driver.find_element.return_value.text = body
    driver.find_elements.return_value = []
    driver.execute_script.return_value = "complete"
    driver.execute_cdp_cmd.side_effect = RuntimeError("no CDP")
    driver.save_screenshot.side_effect = lambda path: Path(path).write_bytes(b"image") or True
    target = tmp_path / "product.png"
    target.write_bytes(b"old evidence")
    def solve(prompt):
        assert not target.exists()
        assert len(list((tmp_path / "challenges").glob("*.json"))) == 1
        driver.find_element.return_value.text = "Raspberry Pi: Add to Cart"
        driver.title = "SC0194(9)"
        return ""
    capture_evidence(driver, target, interactive=True, prompt=solve, settle_timeout=0)
    assert target.read_bytes() == b"image"
    driver.quit.assert_not_called()


@pytest.mark.parametrize("header_row", [1, 2])
def test_share_link_is_saved_to_every_order_row_and_preserves_workbook(tmp_path, header_row):
    path = tmp_path / "synced.xlsx"
    row = workbook(path, header_row)
    wb = openpyxl.load_workbook(path)
    ws = wb["Ordering"]
    ws.append(["old_amazon_order", "38", ITEM["link"], "DigiKey", 100, 5, f"=E{row+2}*F{row+2}", "old-cart"])
    ws.tables["OrderT"].ref = f"A{header_row}:H{row+2}"
    wb.save(path)
    wb.close()
    assert update_order_table_links(str(path), "old_amazon_order", URL, expected_rows=2) == 2
    wb = openpyxl.load_workbook(path)
    ws = wb["Ordering"]
    for r in (row, row + 2):
        assert ws.cell(r, 8).value == URL
        assert ws.cell(r, 8).hyperlink.target == URL
        assert ws.cell(r, 7).value == f"=E{r}*F{r}"
    assert ws.cell(row + 1, 8).value == "keep-cart"
    assert ws.cell(row, 3).font.color.rgb == "000000FF"
    assert ws.tables["OrderT"].ref == f"A{header_row}:H{row+2}"
    assert wb["Bills"].cell(2, 2).value == 10
    wb.close()


@pytest.mark.parametrize("problem", ["missing_column", "wrong_count", "missing_order"])
def test_invalid_link_destination_does_not_claim_success_or_modify_workbook(tmp_path, problem):
    path = tmp_path / "synced.xlsx"
    workbook(path)
    if problem == "missing_column":
        wb = openpyxl.load_workbook(path)
        wb["Ordering"].cell(2, 8).value = "Notes"
        wb.save(path)
        wb.close()
    before = path.read_bytes()
    with pytest.raises(ValueError):
        update_order_table_links(str(path), "missing" if problem == "missing_order" else "old_amazon_order", URL,
                                 expected_rows=2 if problem == "wrong_count" else 1)
    assert path.read_bytes() == before


@pytest.mark.parametrize("cancel", [False, True])
def test_failed_excel_save_retains_cart_and_recovery_url_until_retry_or_cancel(monkeypatch, tmp_path, cancel):
    monkeypatch.chdir(tmp_path)
    driver = MagicMock()
    monkeypatch.setattr(purchase_cart.webdriver, "Chrome", lambda **kw: driver)
    monkeypatch.setattr(purchase_cart, "Service", MagicMock())
    monkeypatch.setattr(purchase_cart, "prepare_digikey_items", MagicMock())
    monkeypatch.setattr(purchase_cart, "navigate_for_evidence", MagicMock())
    monkeypatch.setattr(purchase_cart, "capture_evidence", lambda d, p, **kw: str(p))
    monkeypatch.setattr(purchase_cart, "verify_digikey_cart_with_retry", lambda *a: None)
    monkeypatch.setattr(share_a_cart, "find_share_a_cart_extension", lambda: None)
    monkeypatch.setattr(share_a_cart, "prompt_for_share_a_cart", lambda *a, **kw: URL)
    answers = iter(["100", "500", "y", "", "", "500", "cancel" if cancel else ""])
    def answer(prompt):
        driver.quit.assert_not_called()
        if "retry saving" in prompt:
            recovery = tmp_path / "screenshots" / "test" / "share_a_cart.json"
            assert json.loads(recovery.read_text())["share_url"] == URL
        return next(answers)
    monkeypatch.setattr("builtins.input", answer)
    saver = MagicMock(side_effect=[PermissionError("Excel is open"), None])
    if cancel:
        with pytest.raises(ValueError, match="Engage preparation cancelled"):
            purchase_cart.prepare_cart([ITEM.copy()], "DigiKey", "test", save_share_url=saver)
        driver.quit.assert_called_once()
        assert saver.call_count == 1
    else:
        cart = purchase_cart.prepare_cart([ITEM.copy()], "DigiKey", "test", save_share_url=saver)
        assert cart["share_url"] == URL
        driver.quit.assert_not_called()
        assert saver.call_count == 2


def test_required_share_url_reprompts_and_allows_explicit_cancel(monkeypatch):
    answers = iter(["", "https://www.digikey.com/ordering/shoppingcart", URL, "cancel"])
    monkeypatch.setattr("builtins.input", lambda p: next(answers))
    assert share_a_cart.prompt_for_share_a_cart(1, "DigiKey", required=True) == URL
    assert share_a_cart.prompt_for_share_a_cart(1, "DigiKey", required=True) is None
