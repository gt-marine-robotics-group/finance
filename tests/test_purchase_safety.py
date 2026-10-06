"""Checks for stale quotes, browser challenges, and custom Engage field routing."""

import base64
import json
import os
from decimal import Decimal
from pathlib import Path
from unittest.mock import MagicMock
from urllib.parse import quote

import openpyxl
import pandas as pd
import pytest
from selenium.webdriver.common.by import By
from selenium.common.exceptions import ElementClickInterceptedException, StaleElementReferenceException, TimeoutException

from mrg_finance import cli, order_excel_builder, share_a_cart
from mrg_finance.spreadsheet_utils import get_col_val
from mrg_finance.engage_fields import FormFieldError, fill_purchase_fields, find_labeled_field
from mrg_finance.purchase_validation import money, read_amazon_cart, reconcile_amazon_cart, require_matching_total
from mrg_finance.screenshot_capture import BrowserChallenge, capture_evidence, navigate_for_evidence, save_page_screenshot
from mrg_finance.vendor_payee import lookup_vendor_payee
from mrg_finance.purchase_cart import (
    add_amazon_item_to_cart, decline_amazon_coverage, prepare_cart, prompt_money,
    recheck_cart, verify_amazon_cart_with_retry,
)


class Field:
    def __init__(self, id_, value="", text="", **attributes):
        self.attributes = {"id": id_, "value": value, **attributes}
        self.text = text

    def clear(self):
        self.attributes["value"] = ""

    def send_keys(self, value):
        self.attributes["value"] += str(value)

    def get_attribute(self, name):
        return self.attributes.get(name)

    def is_displayed(self):
        return True


@pytest.mark.parametrize("value,expected", [("$1,234.56", "1234.56"), (0.1 + 0.2, "0.30"), ("2.675", "2.68")])
def test_currency_rounds_to_cents(value, expected):
    assert money(value) == Decimal(expected)


@pytest.mark.parametrize("value", ["nan", "inf", "-1", "", None, "1e999"])
def test_invalid_currency_is_not_a_fallback(value):
    with pytest.raises(ValueError):
        money(value)


def test_one_cent_difference_blocks_engage():
    require_matching_total("19.99", "19.99")
    with pytest.raises(ValueError, match="differs"):
        require_matching_total("19.99", "20.00")


def order():
    return [{"item_name": "Part", "quantity": 2, "cost": 10,
             "link": "https://www.amazon.com/dp/B012345678"}]


def snapshot(quantity=2, price="12.50", extra=False):
    items = [{"asin": "B012345678", "quantity": quantity, "unit_price": Decimal(price)}]
    if extra:
        items.append({"asin": "B000000001", "quantity": 1, "unit_price": Decimal("5")})
    return {"items": items, "subtotal": sum(i["unit_price"] * i["quantity"] for i in items)}


def test_actual_cart_quote_retains_approved_baseline():
    items = order()
    reconcile_amazon_cart(items, snapshot())
    assert items[0]["cost"] == 10
    assert items[0]["quoted_unit_cost"] == 12.5
    assert items[0]["quoted_total"] == 25


@pytest.mark.parametrize("cart", [snapshot(quantity=1), snapshot(extra=True), {"items": [], "subtotal": Decimal(0)}])
def test_seller_limits_and_unrelated_items_block(cart):
    with pytest.raises(ValueError, match="quantities"):
        reconcile_amazon_cart(order(), cart)


def test_duplicate_asins_are_aggregated():
    items = order() + order()
    reconcile_amazon_cart(items, snapshot(quantity=4))
    assert sum(i["quoted_total"] for i in items) == 50


def test_cart_reader_refuses_unreadable_prices():
    driver = MagicMock()
    row = MagicMock()
    row.is_displayed.return_value = True
    row.get_attribute.side_effect = lambda key: {"data-asin": "B012345678", "data-quantity": "2"}.get(key)
    row.find_elements.return_value = []
    driver.find_elements.return_value = [row]
    with pytest.raises(ValueError, match="Could not read"):
        read_amazon_cart(driver)


def amazon_cart_driver(*, modern=True, stepper=False, subtotal="609.90", hidden=False):
    """Use the price/quantity markup observed on the actual Amazon cart."""
    row = MagicMock()
    row.is_displayed.return_value = True
    row.get_attribute.side_effect = lambda name: {
        "data-asin": "B07TC2BK1X", "data-quantity": None if stepper else "5",
    }.get(name)
    price = MagicMock()
    price.is_displayed.return_value = not hidden
    price.text = "$121\n.\n98" if modern else "$121.98"
    price.find_elements.side_effect = lambda by, selector: (
        [Field("accessible-price", textContent="$121.98")] if modern and selector == ".a-offscreen" else []
    )
    quantity_controls = [
        Field("decrease", **{"aria-label": "Decrease quantity by one, Quantity is 5, Raspberry Pi 4"}),
        Field("increase", **{"aria-label": "Increase quantity by one, Quantity is 5, Raspberry Pi 4"}),
    ]

    def row_elements(by, selector):
        if selector == "button[aria-label*='Quantity is']":
            return quantity_controls
        if selector == (".apex-price-to-pay-value" if modern else ".sc-product-price, .sc-price"):
            return [price]
        return []

    row.find_elements.side_effect = row_elements
    total = MagicMock()
    total.is_displayed.return_value = True
    total.text = f"${subtotal}"
    total.find_elements.return_value = []
    driver = MagicMock()
    driver.find_elements.side_effect = lambda by, selector: (
        [row] if selector == "#sc-active-cart .sc-list-item[data-asin]" else
        [total] if "#sc-subtotal-amount-buybox" in selector else []
    )
    return driver


@pytest.mark.parametrize("modern,stepper", [(True, False), (True, True), (False, False)])
def test_cart_reader_reads_current_quote_and_sidebar_subtotal(modern, stepper):
    result = read_amazon_cart(amazon_cart_driver(modern=modern, stepper=stepper))
    assert result == {"items": [{"asin": "B07TC2BK1X", "quantity": 5, "unit_price": Decimal("121.98")}],
                      "subtotal": Decimal("609.90")}


def test_cart_reader_still_rejects_a_one_cent_subtotal_mismatch():
    with pytest.raises(ValueError, match="differs"):
        read_amazon_cart(amazon_cart_driver(subtotal="609.89"))


def test_cart_reader_does_not_accept_hidden_price_nodes():
    with pytest.raises(ValueError, match="missing unit price"):
        read_amazon_cart(amazon_cart_driver(hidden=True))


def test_add_to_cart_waits_for_count_increase(monkeypatch, tmp_path):
    driver = MagicMock()
    button = MagicMock()
    counts = iter(["0", "0", "5"])

    def elements(by, selector):
        if selector == "nav-cart-count":
            return [Field("nav-cart-count", text=next(counts))]
        return []

    driver.find_elements.side_effect = elements

    class Wait:
        def __init__(self, *args, **kwargs):
            pass

        def until(self, predicate):
            button.click.assert_called_once()
            assert predicate(driver) is False  # Amazon has not processed the click yet.
            assert predicate(driver) is True

    monkeypatch.setattr("mrg_finance.purchase_cart.WebDriverWait", Wait)
    add_amazon_item_to_cart(driver, button, "Part", tmp_path)
    driver.get.assert_not_called()


def test_add_to_cart_accepts_explicit_confirmation(monkeypatch, tmp_path):
    driver = MagicMock()
    button = MagicMock()
    driver.find_elements.side_effect = lambda by, selector: (
        [] if selector == "nav-cart-count" else [Field("confirmation", text="Added to cart")]
    )

    class Wait:
        def __init__(self, *args, **kwargs):
            pass

        def until(self, predicate):
            assert predicate(driver) is True

    monkeypatch.setattr("mrg_finance.purchase_cart.WebDriverWait", Wait)
    add_amazon_item_to_cart(driver, button, "Part", tmp_path)
    button.click.assert_called_once()


def test_intercepted_add_reacquires_visible_button_then_confirms_once(monkeypatch, tmp_path):
    driver = MagicMock()
    blocked, hidden, current = MagicMock(), MagicMock(), MagicMock()
    blocked.click.side_effect = ElementClickInterceptedException("offer overlaps button")
    hidden.is_displayed.return_value = False
    count = Field("nav-cart-count", text="0")
    current.click.side_effect = lambda: setattr(count, "text", "5")
    driver.find_elements.side_effect = lambda by, selector: (
        [count] if selector == "nav-cart-count" else [hidden, current]
        if selector == "add-to-cart-button" else []
    )

    class Wait:
        def __init__(self, *args, **kwargs):
            pass

        def until(self, predicate):
            assert predicate(driver) is True

    monkeypatch.setattr("mrg_finance.purchase_cart.WebDriverWait", Wait)
    add_amazon_item_to_cart(driver, blocked, "Part", tmp_path)
    blocked.click.assert_called_once()
    hidden.click.assert_not_called()
    current.click.assert_called_once()
    driver.quit.assert_not_called()


def test_replaced_add_button_does_not_duplicate_already_added_items(monkeypatch, tmp_path):
    driver, button = MagicMock(), MagicMock()
    count = Field("nav-cart-count", text="0")
    driver.find_elements.side_effect = lambda by, selector: [count] if selector == "nav-cart-count" else []

    def replace_after_add():
        count.text = "5"
        raise StaleElementReferenceException("buy box replaced")

    button.click.side_effect = replace_after_add

    class Wait:
        def __init__(self, *args, **kwargs):
            pass

        def until(self, predicate):
            assert predicate(driver) is True

    monkeypatch.setattr("mrg_finance.purchase_cart.WebDriverWait", Wait)
    add_amazon_item_to_cart(driver, button, "Part", tmp_path)
    button.click.assert_called_once()


def test_persistent_click_overlay_allows_manual_completion_in_same_browser(monkeypatch, tmp_path):
    driver, button = MagicMock(), MagicMock()
    driver.find_elements.return_value = []
    button.click.side_effect = ElementClickInterceptedException("offer overlaps button")
    wait = MagicMock()
    wait.until.side_effect = TimeoutException()
    monkeypatch.setattr("mrg_finance.purchase_cart.WebDriverWait", lambda *a, **kw: wait)
    screenshot = MagicMock(return_value=str(tmp_path / "blocked.png"))
    monkeypatch.setattr("mrg_finance.purchase_cart.save_page_screenshot", screenshot)

    def manually_finish(prompt):
        assert "Finish adding" in prompt
        driver.get.assert_not_called()
        driver.quit.assert_not_called()
        return ""

    monkeypatch.setattr("builtins.input", manually_finish)
    add_amazon_item_to_cart(driver, button, "Part", tmp_path)
    screenshot.assert_called_once()
    driver.quit.assert_not_called()


def coverage_control(label, *, label_id=None):
    control = MagicMock()
    control.text = "" if label_id else label
    control.get_attribute.side_effect = lambda key: {"aria-labelledby": label_id}.get(key)
    return control


@pytest.mark.parametrize("before_click", [False, True])
def test_coverage_offer_declines_once_before_confirming_add(monkeypatch, tmp_path, before_click):
    driver, button, dialog = MagicMock(), MagicMock(), MagicMock()
    decline = coverage_control("No Thanks", label_id="attachSiNoCoverage-announce")
    accept = coverage_control("Add Protection")
    dialog.text = "Add a protection plan? Coverage for accidents."
    dialog.is_displayed.return_value = before_click
    dialog.find_elements.side_effect = lambda by, selector: (
        [Field("attachSiNoCoverage-announce", textContent="No Thanks")] if by == By.ID else [accept, decline]
    )
    count = Field("nav-cart-count", text="0")
    driver.find_elements.side_effect = lambda by, selector: (
        [count] if selector == "nav-cart-count" else [button]
        if selector == "add-to-cart-button" else [dialog] if "[role='dialog']" in selector else []
    )

    def click_product():
        dialog.is_displayed.return_value = True
        if before_click:
            raise ElementClickInterceptedException("coverage dialog intercepts buy box")
        count.text = "5"  # Amazon may update the count before the coverage dialog closes.

    button.click.side_effect = click_product
    decline.click.side_effect = lambda: setattr(count, "text", "5")
    waits = []

    class Wait:
        def __init__(self, d, timeout, **kwargs):
            waits.append(timeout)

        def until(self, predicate):
            if len(waits) == 1:
                assert predicate(driver) is False
                decline.click.assert_called_once()
                assert predicate(driver) is False  # Still loading; don't click decline again.
                decline.click.assert_called_once()
                dialog.is_displayed.return_value = False
            assert predicate(driver) is True

    monkeypatch.setattr("mrg_finance.purchase_cart.WebDriverWait", Wait)
    add_amazon_item_to_cart(driver, button, "Part", tmp_path)
    button.click.assert_called_once()
    decline.click.assert_called_once()
    accept.click.assert_not_called()
    driver.quit.assert_not_called()


@pytest.mark.parametrize("text,labels,visible", [
    ("Protection coverage", ["Add Protection"], True),
    ("Protection coverage", ["No Thanks", "Decline coverage"], True),
    ("Try Prime", ["No Thanks"], False),
])
def test_coverage_handler_leaves_ambiguous_or_unrelated_controls_alone(text, labels, visible):
    driver, dialog = MagicMock(), MagicMock()
    controls = [coverage_control(label) for label in labels]
    dialog.text = text
    dialog.find_elements.return_value = controls
    driver.find_elements.return_value = [dialog]
    assert decline_amazon_coverage(driver, set()) is visible
    for control in controls:
        control.click.assert_not_called()


def test_generic_popup_dismissal_does_not_click_numbered_amazon_offer_button(monkeypatch):
    driver, offer = MagicMock(), MagicMock()
    driver.page_source = "Add Protection"
    driver.find_elements.side_effect = lambda by, selector: [offer] if selector == "#a-autoid-0-announce" else []
    from mrg_finance.price_scraper import dismiss_popups_and_interstitials
    dismiss_popups_and_interstitials(driver)
    offer.click.assert_not_called()


@pytest.mark.parametrize("answers,default,expected", [
    (["", "oops", "-1", "nan", "$609.90"], None, "609.90"),
    (["oops", ""], "0", "0.00"),
])
def test_money_prompt_recovers_from_blank_and_invalid_entries(monkeypatch, capsys, answers, default, expected):
    remaining = iter(answers)
    monkeypatch.setattr("builtins.input", lambda p: next(remaining))
    assert prompt_money("Amount: $", default=default) == Decimal(expected)
    assert list(remaining) == []
    assert "browser is still open" in capsys.readouterr().out


def test_money_prompt_can_cancel_after_bad_input(monkeypatch):
    remaining = iter(["", "cancel"])
    monkeypatch.setattr("builtins.input", lambda p: next(remaining))
    with pytest.raises(ValueError, match="entry cancelled"):
        prompt_money("Amount: $")


def test_unconfirmed_add_keeps_browser_open_for_manual_completion(monkeypatch, tmp_path):
    driver = MagicMock()
    driver.find_elements.return_value = []
    wait = MagicMock()
    wait.until.side_effect = TimeoutException()
    monkeypatch.setattr("mrg_finance.purchase_cart.WebDriverWait", lambda *a, **kw: wait)
    monkeypatch.setattr("mrg_finance.purchase_cart.save_page_screenshot", lambda d, p, **kw: str(p))

    def cancel(prompt):
        assert "Finish adding" in prompt
        driver.get.assert_not_called()
        driver.quit.assert_not_called()
        return "cancel"

    monkeypatch.setattr("builtins.input", cancel)
    with pytest.raises(ValueError, match="cancelled"):
        add_amazon_item_to_cart(driver, MagicMock(), "Part", tmp_path)


@pytest.mark.parametrize("error", [ValueError("empty cart"), StaleElementReferenceException("loading")])
def test_initial_cart_verification_retries_in_same_browser(monkeypatch, tmp_path, error):
    driver = MagicMock()
    driver.current_url = "https://www.amazon.com/gp/cart/view.html"
    driver.title = "Amazon cart"
    reader = MagicMock(side_effect=[error, snapshot(quantity=1), snapshot()])
    monkeypatch.setattr("mrg_finance.purchase_cart.read_amazon_cart", reader)

    def capture(d, path, **kwargs):
        assert d is driver
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(b"test screenshot")
        return str(path)

    monkeypatch.setattr("mrg_finance.purchase_cart.capture_evidence", capture)
    monkeypatch.setattr("mrg_finance.purchase_cart.save_page_screenshot", capture)
    prompts = []

    def correct(prompt):
        prompts.append(prompt)
        assert "retry verification" in prompt
        assert not (tmp_path / "cart.png").exists()
        driver.quit.assert_not_called()
        return ""

    monkeypatch.setattr("builtins.input", correct)
    items = order()
    result = verify_amazon_cart_with_retry(driver, items, tmp_path)
    assert result == snapshot()
    assert items[0]["quoted_total"] == 25
    assert len(prompts) == 2
    assert (tmp_path / "cart.png").exists()
    assert len(list((tmp_path / "cart_diagnostics").glob("*.json"))) == 2
    driver.quit.assert_not_called()


def test_cancelled_verification_does_not_produce_quote(monkeypatch, tmp_path):
    driver = MagicMock()
    driver.current_url = "https://www.amazon.com/gp/cart/view.html"
    driver.title = "Empty cart"
    monkeypatch.setattr("mrg_finance.purchase_cart.read_amazon_cart", MagicMock(side_effect=ValueError("empty cart")))

    def capture(d, path, **kwargs):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(b"test screenshot")
        return str(path)

    monkeypatch.setattr("mrg_finance.purchase_cart.capture_evidence", capture)
    monkeypatch.setattr("mrg_finance.purchase_cart.save_page_screenshot", capture)
    monkeypatch.setattr("builtins.input", lambda prompt: "cancel")
    items = order()
    with pytest.raises(ValueError, match="verification cancelled"):
        verify_amazon_cart_with_retry(driver, items, tmp_path)
    assert not (tmp_path / "cart.png").exists()
    assert "quoted_unit_cost" not in items[0]


def test_refresh_rejects_changed_unit_prices_even_when_total_is_unchanged(monkeypatch, tmp_path):
    old = snapshot()
    old["items"].append({"asin": "B000000001", "quantity": 1, "unit_price": Decimal("5")})
    old["subtotal"] = Decimal("30")
    fresh = snapshot(price="13")
    fresh["items"].append({"asin": "B000000001", "quantity": 1, "unit_price": Decimal("4")})
    fresh["subtotal"] = Decimal("30")
    monkeypatch.setattr("mrg_finance.purchase_cart.navigate_for_evidence", lambda *a: None)
    monkeypatch.setattr("mrg_finance.purchase_cart.capture_evidence", lambda *a, **kw: None)
    monkeypatch.setattr("mrg_finance.purchase_cart.read_amazon_cart", lambda d: fresh)
    cart = {"driver": MagicMock(), "vendor_key": "amazon", "snapshot": old,
            "subtotal": Decimal("30"), "screenshot": str(tmp_path / "cart.png")}
    with pytest.raises(ValueError, match="changed after review"):
        recheck_cart(cart, order())


@pytest.mark.parametrize("total,changed", [("25.00", False), ("25.01", True)])
def test_preupload_total_retries_typing_errors_but_still_blocks_changed_quote(monkeypatch, tmp_path, total, changed):
    driver = MagicMock()
    monkeypatch.setattr("mrg_finance.purchase_cart.navigate_for_evidence", lambda *a: None)
    monkeypatch.setattr("mrg_finance.purchase_cart.capture_evidence", lambda *a, **kw: None)
    monkeypatch.setattr("mrg_finance.purchase_cart.read_amazon_cart", lambda d: snapshot())
    cart = {"driver": driver, "vendor_key": "amazon", "snapshot": snapshot(),
            "subtotal": Decimal("25"), "total": Decimal("25"), "screenshot": str(tmp_path / "cart.png")}
    remaining = iter(["", "oops", total])

    def answer(prompt):
        driver.quit.assert_not_called()
        return next(remaining)

    monkeypatch.setattr("builtins.input", answer)
    if changed:
        with pytest.raises(ValueError, match="differs"):
            recheck_cart(cart, order())
    else:
        recheck_cart(cart, order())
    assert list(remaining) == []


def challenge_driver(body):
    driver = MagicMock()
    driver.current_url = "https://www.amazon.com/errors/validateCaptcha"
    driver.title = "Robot Check"
    driver.find_element.return_value.text = body
    driver.find_elements.return_value = []
    driver.execute_script.return_value = "complete"
    driver.save_screenshot.side_effect = lambda path: Path(path).write_bytes(b"image") or True
    return driver


def test_captcha_saves_diagnostic_and_invalidates_stale_attachment(tmp_path):
    target = tmp_path / "cart.png"
    target.write_bytes(b"old quote")
    with pytest.raises(BrowserChallenge):
        capture_evidence(challenge_driver("Enter the characters you see below"), target)
    assert not target.exists()
    metadata = list((tmp_path / "challenges").glob("*.json"))
    assert len(metadata) == 1
    assert json.loads(metadata[0].read_text())["status"] == "challenge"


def test_challenge_can_be_solved_before_capturing_quote(tmp_path):
    driver = challenge_driver("Robot check")
    def solve(_):
        driver.current_url = "https://www.amazon.com/gp/cart/view.html"
        driver.title = "Cart"
        driver.find_element.return_value.text = "Subtotal $25.00"
        return ""
    driver.execute_cdp_cmd.side_effect = RuntimeError("No CDP")
    path = capture_evidence(driver, tmp_path / "cart.png", interactive=True, prompt=solve)
    assert Path(path).exists()
    assert list((tmp_path / "challenges").glob("*.png"))


def test_full_page_capture_and_viewport_fallback(tmp_path):
    driver = MagicMock()
    driver.execute_cdp_cmd.side_effect = [{"contentSize": {"width": 100, "height": 5000}},
                                         {"data": base64.b64encode(b"png bytes").decode()}]
    target = tmp_path / "page.png"
    save_page_screenshot(driver, target)
    assert target.read_bytes() == b"png bytes"
    assert driver.execute_cdp_cmd.call_args.args[1]["clip"]["height"] == 5000
    driver.execute_cdp_cmd.side_effect = RuntimeError("Unsupported")
    driver.save_screenshot.return_value = False
    with pytest.raises(RuntimeError, match="capture failed"):
        save_page_screenshot(driver, target)


def test_navigation_timeout_allows_challenge_inspection():
    from selenium.common.exceptions import TimeoutException
    driver = MagicMock()
    driver.get.side_effect = TimeoutException()
    navigate_for_evidence(driver, "https://example.org")
    driver.get.assert_called_once()


def form_driver():
    driver = MagicMock()
    amount = Field("Amount")
    driver.find_elements.side_effect = lambda by, value: [amount] if value == "Amount" else []
    fields = {"refs": Field("refs"), "sga": Field("sga"), "payee": Field("payee"), "email": Field("email")}
    def resolve(script, labels, excluded):
        if labels[0].startswith("What is"):
            return [fields["refs"]]
        if labels[0].startswith("Include Bill") or labels[0] == "SGA Budget":
            return [fields["sga"]]
        if labels[0] == "payee name":
            return [fields["payee"]]
        if labels[0] == "payee email":
            return [fields["email"]]
        return []
    driver.execute_script.side_effect = resolve
    return driver, amount, fields


def test_bill_details_go_into_separate_verified_fields():
    driver, amount, fields = form_driver()
    fill_purchase_fields(driver, amount="25", cart_total="25", bill_refs="Bill 376582, Line 4",
                         sga_lines="$25.00, Bill 376582, Line 4", payee={"name": "Amazon"})
    assert fields["refs"].get_attribute("value") == "Bill 376582, Line 4"
    assert fields["sga"].get_attribute("value").startswith("$25.00")
    assert fields["payee"].get_attribute("value") == "Amazon"
    assert amount.get_attribute("value") == "25.00"
    assert all(call.args[1] != "Description" for call in driver.find_elements.call_args_list)


def test_missing_custom_field_stops_without_description_fallback():
    driver, _, _ = form_driver()
    driver.execute_script.return_value = []
    driver.execute_script.side_effect = None
    with pytest.raises(FormFieldError, match="found 0"):
        fill_purchase_fields(driver, amount=25, cart_total=25, bill_refs="Bill 1, Line 2",
                             sga_lines="$25, Bill 1", payee={"name": "Vendor"})
    assert not driver.find_element.called


def test_ambiguous_field_is_rejected():
    driver = MagicMock()
    driver.execute_script.return_value = [Field("a"), Field("b")]
    with pytest.raises(FormFieldError, match="found 2"):
        find_labeled_field(driver, ["SGA Bill"])


def test_requested_amount_mismatch_blocks_before_custom_fields():
    driver, _, _ = form_driver()
    with pytest.raises(ValueError, match="differs"):
        fill_purchase_fields(driver, amount=20, cart_total=25, bill_refs="Bill 1",
                             sga_lines="$20", payee={"name": "Vendor"})
    driver.execute_script.assert_not_called()


@pytest.mark.parametrize("vendor", ["Digi-Key", "DigiKey Electronics", "Other vendor"])
def test_non_amazon_share_cart_never_posts_synthetic_payload(monkeypatch, vendor):
    post = MagicMock()
    monkeypatch.setattr("requests.post", post)
    assert share_a_cart.create_share_a_cart_link(order(), vendor) is None
    post.assert_not_called()


@pytest.mark.parametrize("raw,expected", [
    ("ABC123", "https://share-a-cart.com/get/ABC123"),
    ("share-a-cart.com/get/ABC123", "https://share-a-cart.com/get/ABC123"),
    ("https://shareacart.net/get/ABC123", "https://shareacart.net/get/ABC123"),
    ("https://evil.example/get/ABC123", None), ("not a cart", None),
])
def test_shared_cart_links_are_validated(raw, expected):
    assert share_a_cart.normalize_share_a_cart_url(raw) == expected


def test_known_payees_require_no_options_or_network(monkeypatch):
    get = MagicMock()
    monkeypatch.setattr("requests.get", get)
    assert lookup_vendor_payee("Amazon", "")["name"] == "Amazon.com Services LLC"
    assert lookup_vendor_payee("Digi-Key", "")["email"] == "orders@digikey.com"
    get.assert_not_called()


def test_other_payee_uses_official_structured_contact(monkeypatch):
    response = MagicMock()
    response.url = "https://example.org/contact"
    response.text = '<script type="application/ld+json">' + json.dumps({
        "@type": "Organization", "name": "Example Parts", "url": "https://example.org",
        "address": {"streetAddress": "123 Main St", "addressLocality": "Atlanta", "postalCode": "30332"},
        "email": "orders@example.org",
    }) + '</script>'
    monkeypatch.setattr("requests.get", lambda *a, **kw: response)
    result = lookup_vendor_payee("Example", "https://example.org/products/part")
    assert result["address"] == "123 Main St, Atlanta, 30332"
    assert result["source"] == response.url


def test_report_uses_verified_line_and_quote_instead_of_excel_id(tmp_path):
    item = order()[0]
    item.update(bill_no="376582", bill_item_id="999", resolved_line_id=4,
                resolved_section="B06", quoted_unit_cost=12.5)
    xlsx, _ = order_excel_builder.generate_order_budget_vs_quoted_excel("test", [item], output_dir=str(tmp_path))
    wb = openpyxl.load_workbook(xlsx)
    row = list(wb.active.iter_rows(min_row=5, max_row=5, values_only=True))[0]
    assert row[0] == "Line 4"
    assert row[3] == "B06"
    assert row[5] == 10
    assert row[8] == 12.5
    wb.close()


def test_cli_removed_gui_and_exposes_personal_cart(monkeypatch, capsys):
    monkeypatch.setattr("sys.argv", ["mrg-finance", "purchase", "--help"])
    with pytest.raises(SystemExit) as error:
        cli.main()
    assert error.value.code == 0
    help_text = capsys.readouterr().out
    assert "--cart-source" in help_text
    assert "payee" not in help_text
    assert "review GUI" not in help_text
    assert not hasattr(cli, "cmd_review")


@pytest.mark.parametrize("flag,select_order,vendor,answers,expected", [
    (None, False, "Amazon", [], "automated"),
    ("personal", False, "Amazon", [], "personal"),
    ("automated", False, "Amazon", [], "automated"),
    (None, True, "Amazon", ["1"], "automated"),
    (None, False, "DigiKey", [], "automated"),
    ("personal", False, "DigiKey", [], "personal"),
])
def test_purchase_defaults_to_automated_without_a_source_prompt(
    monkeypatch, capsys, flag, select_order, vendor, answers, expected,
):
    monkeypatch.setattr(cli, "load_ordering", lambda: pd.DataFrame([{
        "Order ID": "test-order", "Bill Item ID": "1", "Quantity": 2, "Status": "pending",
    }]))
    monkeypatch.setattr(cli, "load_xlsx", lambda: pd.DataFrame([{
        "Bill Item ID": "1", "Item": "Part", "Vendor": vendor, "Cost": 10,
    }]))
    monkeypatch.setattr(cli, "get_xlsx_path", lambda: "/test/workbook.xlsx")
    run = MagicMock(return_value=MagicMock(returncode=0))
    monkeypatch.setattr(cli.subprocess, "run", run)
    argv = ["mrg-finance", "purchase"]
    if not select_order:
        argv.extend(["--order", "test-order"])
    if flag:
        argv.extend(["--cart-source", flag])
    monkeypatch.setattr("sys.argv", argv)
    prompts = []
    remaining = iter(answers)

    def answer(prompt):
        prompts.append(prompt)
        return next(remaining)

    monkeypatch.setattr("builtins.input", answer)
    with pytest.raises(SystemExit) as error:
        cli.main()
    assert error.value.code == 0
    run.assert_called_once()
    command = run.call_args.args[0]
    assert command[command.index("--cart-source") + 1] == expected
    assert command[command.index("--order") + 1] == "test-order"
    assert not any("cart source" in p.lower() for p in prompts)
    assert not any("verify" in p.lower() for p in prompts)
    assert any("Select order" in p for p in prompts) is select_order
    output = capsys.readouterr().out
    assert ("Pending Orders" in output) is select_order
    assert list(remaining) == []


def test_personal_mode_shares_real_cart_and_uses_account_profile(monkeypatch, tmp_path, capsys):
    monkeypatch.chdir(tmp_path)
    driver = MagicMock()
    factory = MagicMock(return_value=driver)
    monkeypatch.setattr("mrg_finance.purchase_cart.webdriver.Chrome", factory)
    monkeypatch.setattr("mrg_finance.purchase_cart.Service", MagicMock())
    monkeypatch.setattr("mrg_finance.purchase_cart.navigate_for_evidence", lambda *a: None)
    monkeypatch.setattr("mrg_finance.purchase_cart.capture_evidence", lambda d, p, **kw: str(p))
    monkeypatch.setattr("mrg_finance.purchase_cart.read_amazon_cart", lambda d: snapshot())
    monkeypatch.setattr(share_a_cart, "find_share_a_cart_extension", lambda: None)
    api = MagicMock()
    monkeypatch.setattr(share_a_cart, "create_share_a_cart_link", api)
    monkeypatch.setattr(share_a_cart, "prompt_for_share_a_cart", lambda *a, **kw: "https://share-a-cart.com/get/REAL01")
    # Enter at the required total used to tear down this browser and lose the cart.
    # Correct a malformed fee and a mismatching total without restarting preparation.
    answers = iter(["", "oops", "", "", "", "26.00", "1.00", "", "26.00"])

    def answer(prompt):
        driver.quit.assert_not_called()
        return next(answers)

    monkeypatch.setattr("builtins.input", answer)
    cart = prepare_cart(order(), "Amazon", "personal-test", source="personal")
    assert cart["total"] == Decimal("26.00")
    assert cart["shipping"] == Decimal("1.00")
    assert cart["share_url"].endswith("REAL01")
    assert list(answers) == []
    assert "Check the vendor's charges and re-enter them" in capsys.readouterr().out
    assert any(a.startswith("--user-data-dir=") for a in factory.call_args.kwargs["options"].arguments)
    api.assert_not_called()


@pytest.mark.parametrize("answers,unit_price,shipping,total", [
    (["100", "500", "y", "0", "0", "500"], "100", "0", "500"),
    (["100", "508", "", "500", "y", "8", "0", "508"], "100", "8", "508"),
    (["100", "508", "prices", "101.60", "508", "y", "0", "0", "508"], "101.60", "0", "508"),
    (["100", "508", "cancel"], None, None, None),
])
def test_digikey_default_source_prepares_cart_then_requires_manual_quote_verification(
    monkeypatch, tmp_path, capsys, answers, unit_price, shipping, total,
):
    monkeypatch.chdir(tmp_path)
    driver = MagicMock()
    monkeypatch.setattr("mrg_finance.purchase_cart.webdriver.Chrome", lambda **kwargs: driver)
    monkeypatch.setattr("mrg_finance.purchase_cart.Service", MagicMock())
    navigate = MagicMock()
    monkeypatch.setattr("mrg_finance.purchase_cart.navigate_for_evidence", navigate)
    monkeypatch.setattr("mrg_finance.purchase_cart.capture_evidence", lambda d, p, **kw: str(p))
    # Manual entry is now reached only through the explicitly selected fallback.
    monkeypatch.setattr("mrg_finance.purchase_cart.verify_digikey_cart_with_retry", lambda *a: None)
    prepare = MagicMock()
    monkeypatch.setattr("mrg_finance.purchase_cart.prepare_digikey_items", prepare)
    monkeypatch.setattr(share_a_cart, "find_share_a_cart_extension", lambda: None)
    api = MagicMock()
    monkeypatch.setattr(share_a_cart, "create_share_a_cart_link", api)
    sharing = MagicMock(return_value="https://share-a-cart.com/get/DIGI01")
    monkeypatch.setattr(share_a_cart, "prompt_for_share_a_cart", sharing)
    answers = iter(answers)
    def answer(prompt=""):
        driver.quit.assert_not_called()
        return next(answers)
    monkeypatch.setattr("builtins.input", answer)
    items = order()
    items[0]["quantity"] = 5
    items[0]["link"] = "https://www.digikey.com/en/products/detail/example/part/123"
    if total is None:
        with pytest.raises(ValueError, match="subtotal verification cancelled"):
            prepare_cart(items, "DigiKey", "digikey-test")
        driver.quit.assert_called_once()
        sharing.assert_not_called()
        assert list(answers) == []
        return
    cart = prepare_cart(items, "DigiKey", "digikey-test")
    assert cart["vendor_key"] == "digikey"
    assert cart["total"] == money(total) - money(shipping)
    assert cart["shipping"] == 0
    assert cart["vendor_total"] == money(total)
    assert cart["vendor_shipping"] == money(shipping)
    assert cart["subtotal"] == money(unit_price) * 5
    assert cart["snapshot"] is None
    assert money(items[0]["quoted_unit_cost"]) == money(unit_price)
    assert list(answers) == []
    output = capsys.readouterr().out
    assert "Engage requested amount" not in output
    assert "5 x $100.00 = $500.00" in output
    prepare.assert_called_once_with(driver, items, cart["folder"])
    navigate.assert_called_once_with(driver, "https://www.digikey.com/ordering/shoppingcart")
    api.assert_not_called()


def test_automated_quantity_limit_stops_before_sharing_or_upload(monkeypatch, tmp_path):
    monkeypatch.chdir(tmp_path)
    driver = MagicMock()
    driver.current_url = order()[0]["link"]
    monkeypatch.setattr("mrg_finance.purchase_cart.webdriver.Chrome", lambda **kw: driver)
    monkeypatch.setattr("mrg_finance.purchase_cart.Service", MagicMock())
    monkeypatch.setattr("mrg_finance.purchase_cart.navigate_for_evidence", lambda *a: None)
    monkeypatch.setattr("mrg_finance.purchase_cart.capture_evidence", lambda d, p, **kw: str(p))
    monkeypatch.setattr("mrg_finance.price_scraper.dismiss_popups_and_interstitials", lambda d: None)
    monkeypatch.setattr("mrg_finance.price_scraper.check_and_set_amazon_quantity", lambda *a: (1, True, "seller maximum"))
    monkeypatch.setattr("mrg_finance.purchase_cart.amazon_items_to_add", lambda d, items, folder: items)
    monkeypatch.setattr(share_a_cart, "find_share_a_cart_extension", lambda: None)
    api = MagicMock()
    monkeypatch.setattr(share_a_cart, "create_share_a_cart_link", api)
    with pytest.raises(ValueError, match="available 1"):
        prepare_cart(order(), "Amazon", "limited-test")
    driver.quit.assert_called_once()
    api.assert_not_called()


@pytest.mark.parametrize("column,key", [("Bill Item ID", "item_name"), ("Total Cost", "cost"),
                                       ("Share-A-Cart Link", "link"), ("Bill No. (Engage)", "bill_item_id")])
def test_column_aliases_do_not_take_values_from_other_known_fields(column, key):
    assert get_col_val({column: "wrong value"}, key) == ""


@pytest.mark.parametrize("changed_amount", [False, True])
@pytest.mark.parametrize("replace_vendor", [False, True])
@pytest.mark.parametrize("funding_kind", ["bill", "budget"])
def test_purchase_upload_gate_uses_real_cart_amount(monkeypatch, tmp_path, capsys, changed_amount, replace_vendor, funding_kind):
    from mrg_finance import automation_purchase
    from mrg_finance.purchase_sources import CartReplacement
    replacement_url = "https://www.digikey.com/en/products/detail/example/part/123"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bills"
    ws.append(["Bill Item ID", "Bill No.", "Bill Title", "Item Name", "Vendor", "Cost", "Quantity", "Link"])
    ws.append(["999", "376582", "Test Bill", "Part", "Amazon", 10, 2, order()[0]["link"]])
    ws = wb.create_sheet("Ordering")
    ws.append(["Order ID", "Bill Item ID", "Quantity", "Status", "Link", "Vendor", "Share-A-Cart Link", "Engage Request Link", "Cost"])
    # Ordering's current price must not replace the approved Bills baseline.
    ws.append(["purchase-test", "999", 2, "Pending", order()[0]["link"], "Amazon", "", "", 40])
    path = tmp_path / "test.xlsx"
    wb.save(path)
    wb.close()
    items = order()
    reconcile_amazon_cart(items, snapshot())
    folder = tmp_path / "evidence"
    folder.mkdir()
    screenshot = folder / "cart.png"
    screenshot.write_bytes(b"fake cart screenshot")
    cart_driver = MagicMock()
    cart = {"driver": cart_driver, "folder": folder, "screenshot": str(screenshot),
            "total": Decimal("25.00"), "subtotal": Decimal("25.00"),
            "shipping": Decimal(0), "tax": Decimal(0), "share_url": "https://share-a-cart.com/get/REAL01"}
    prior_cart_driver = MagicMock()
    prepared_vendors = []
    def prepared(requests, vendor, *a, **kw):
        prepared_vendors.append(vendor)
        quoted_unit = 8.5 if vendor == "DigiKey" else 12.5
        requests[0].update(quoted_unit_cost=quoted_unit, quoted_total=quoted_unit * 2)
        try:
            kw["quote_review"](requests, vendor)
        except CartReplacement:
            prior_cart_driver.quit()
            raise
        cart["subtotal"] = cart["total"] = money(quoted_unit * 2)
        if vendor == "DigiKey":
            from mrg_finance.digikey_quote import procurement_quote
            cart["shipping"] = money("8.49")
            cart["total"] += cart["shipping"]
            cart.update(procurement_quote(cart))
        kw["save_share_url"](cart["share_url"])
        return cart
    monkeypatch.setattr(automation_purchase, "prepare_cart", prepared)
    monkeypatch.setattr(automation_purchase, "USERNAME", "")
    monkeypatch.setattr(automation_purchase, "PASSWORD", "")
    monkeypatch.delenv("ENGAGE_USERNAME", raising=False)
    monkeypatch.delenv("ENGAGE_PASSWORD", raising=False)
    monkeypatch.setattr("sys.argv", ["purchase", "--order", "purchase-test", "--excel-path", str(path)])
    driver, amount, fields = form_driver()
    description = Field("Description")
    uploads = [Field("upload-1"), Field("upload-2")]
    driver.current_url = "https://gatech.campuslabs.com/engage/finance/request/1234"
    driver.find_elements.side_effect = lambda by, selector: [amount] if selector == "Amount" else uploads if selector == 'input[type="file"]' else []
    def find_element(by, selector):
        if selector == "discovery-bar":
            raise LookupError()
        if selector == "Description":
            return description
        return MagicMock()
    driver.find_element.side_effect = find_element
    monkeypatch.setattr(automation_purchase.webdriver, "Chrome", lambda **kw: driver)
    monkeypatch.setattr(automation_purchase.WebDriverWait, "until", lambda *a: MagicMock())
    monkeypatch.setattr(automation_purchase.time, "sleep", lambda *a: None)
    monkeypatch.setattr(automation_purchase, "lookup_bill_item_locations", lambda *a: {"Part": {"section_line_number": 4, "section": "B06", "funding_kind": funding_kind}})
    def recheck(*a):
        if changed_amount:
            amount.attributes["value"] = "24.99"
    monkeypatch.setattr(automation_purchase, "recheck_cart", recheck)
    monkeypatch.setattr(automation_purchase, "save_page_screenshot", lambda *a: "diagnostic.png")
    monkeypatch.setattr("subprocess.run", lambda *a, **kw: MagicMock(returncode=0))
    answers = iter(["y", "n", replacement_url if replace_vendor else "", "", "test-only", "", "", "cancel" if changed_amount else "", ""])
    def answer(prompt=""):
        if "GT username" in prompt:
            assert cart["total"] == money("17.00" if replace_vendor else "25.00")
            assert prepared_vendors == (["Amazon", "DigiKey"] if replace_vendor else ["Amazon"])
            preview = openpyxl.load_workbook(folder / "Budget_vs_Quoted_Detail_purchase-test.xlsx")
            assert preview.active["I5"].value == (8.5 if replace_vendor else 12.5)
            assert preview.active["A5"].value == "Pending lookup"
            assert "Cart Reconciliation" in preview.sheetnames
            preview.close()
        if "Review the comparison spreadsheet" in prompt:
            report = openpyxl.load_workbook(folder / "Budget_vs_Quoted_Detail_purchase-test.xlsx")
            reconciliation = dict(report["Cart Reconciliation"].values)
            assert reconciliation["Requested amount"] == (17 if replace_vendor else 25)
            assert reconciliation["Quote status"] == "Funding references resolved"
            report.close()
        return next(answers)
    monkeypatch.setattr("builtins.input", answer)
    monkeypatch.setattr(automation_purchase.getpass, "getpass", lambda prompt: "test-only")
    if changed_amount:
        with pytest.raises(SystemExit) as error:
            automation_purchase.main()
        assert "cancelled" in str(error.value.code)
        assert all(not f.get_attribute("value") for f in uploads)
    else:
        automation_purchase.main()
        assert amount.get_attribute("value") == ("17.00" if replace_vendor else "25.00")
        assert all(f.get_attribute("value") for f in uploads)
        expected_total = "17.00" if replace_vendor else "25.00"
        assert fields["sga"].get_attribute("value") == f"${expected_total}, Line 4, {funding_kind.title()} 376582, B06"
        report = openpyxl.load_workbook(folder / "Budget_vs_Quoted_Detail_purchase-test.xlsx")
        assert "Cart Reconciliation" in report.sheetnames
        assert report.sheetnames.count("Cart Reconciliation") == 1
        if replace_vendor:
            reconciliation = dict(report["Cart Reconciliation"].values)
            assert reconciliation["Shipping"] == 0
            assert reconciliation["Vendor displayed shipping"] == 8.49
            assert reconciliation["Vendor total"] == 25.49
            assert reconciliation["Engage amount"] == 17
        report.close()
    assert "Share-A-Cart" in description.get_attribute("value")
    assert "Bill 376582" not in description.get_attribute("value")
    cart_driver.quit.assert_called_once()
    driver.quit.assert_called_once()
    saved = openpyxl.load_workbook(path)
    assert saved["Ordering"].cell(2, 7).value == cart["share_url"]
    assert saved["Ordering"].cell(2, 7).hyperlink.target == cart["share_url"]
    saved.close()
    if replace_vendor:
        assert prepared_vendors == ["Amazon", "DigiKey"]
        prior_cart_driver.quit.assert_called_once()
        assert fields["payee"].get_attribute("value") == "Digi-Key Electronics"
        assert "Marine Robotics Group DigiKey Purchase Request" in capsys.readouterr().out
        wb = openpyxl.load_workbook(path)
        assert wb["Ordering"].cell(2, 5).value == replacement_url
        assert wb["Ordering"].cell(2, 6).value == "DigiKey"
        assert wb["Bills"].cell(2, 8).value == order()[0]["link"]
        wb.close()
    else:
        assert prepared_vendors == ["Amazon"]


@pytest.mark.skipif(os.environ.get("MRG_RUN_BROWSER_TESTS") != "1", reason="Opt-in local headless Chrome DOM verification")
def test_real_dom_routes_custom_questions_without_touching_description(monkeypatch, tmp_path):
    from selenium import webdriver
    from selenium.webdriver.chrome.options import Options
    options = Options()
    options.add_argument("--headless=new")
    options.add_argument("--no-sandbox")
    driver = webdriver.Chrome(options=options)
    try:
        html = '''<form>
          <label for="Subject">Subject</label><input id="Subject">
          <label for="Description">Description</label><textarea id="Description">Keep shared cart link</textarea>
          <label for="Amount">Requested Amount</label><input id="Amount">
          <div><p>What is the Budget/Bill # and Request Line #? (Ex. Bill 376582, Line 4)</p><textarea id="custom-1"></textarea></div>
          <div><h3>SGA Bill</h3><label for="custom-2">Include Bill # and total reimbursement amount below ($ Per line item)</label><textarea id="custom-2"></textarea></div>
          <span id="payee-label">Payee Name</span><input id="custom-3" aria-labelledby="payee-label">
          <label for="address">Payee Address</label><textarea id="address"></textarea>
        </form>'''
        driver.get("data:text/html," + quote(html))
        fill_purchase_fields(driver, amount=25, cart_total=25,
                             bill_refs="Bill 376582, Line 4", sga_lines="$25.00, Bill 376582, Line 4",
                             payee={"name": "Amazon.com Services LLC", "address": "410 Terry Avenue North"})
        assert driver.find_element(By.ID, "Description").get_attribute("value") == "Keep shared cart link"
        assert driver.find_element(By.ID, "custom-1").get_attribute("value") == "Bill 376582, Line 4"
        assert driver.find_element(By.ID, "custom-2").get_attribute("value").startswith("$25.00")
        assert driver.find_element(By.ID, "custom-3").get_attribute("value") == "Amazon.com Services LLC"
        assert driver.find_element(By.ID, "address").get_attribute("value") == "410 Terry Avenue North"
        # Current Amazon markup: a split visual price, an accessible price span,
        # a sidebar-only subtotal, and unrelated recommendation prices outside the cart.
        for modern in (True, False):
            price = ('''<span class="a-price apex-price-to-pay-value">
                <span class="a-offscreen" style="position:absolute;left:-10000px">$121.98</span>
                <span aria-hidden="true"><span class="a-price-symbol">$</span>
                <span class="a-price-whole">121<span class="a-price-decimal">.</span></span>
                <span class="a-price-fraction">98</span></span></span>''' if modern else
                '<span class="sc-product-price">$121.98</span>')
            subtotal_id = "sc-subtotal-amount-buybox" if modern else "sc-subtotal-amount-activecart"
            html = f'''<div id="sc-active-cart">
                <div class="sc-list-item" data-asin="B07TC2BK1X" data-quantity="5">{price}</div>
                </div><span id="{subtotal_id}">$609.90</span>
                <div class="recommendations"><span class="sc-product-price">$99.99</span></div>'''
            driver.get("data:text/html," + quote(html))
            quote_snapshot = read_amazon_cart(driver)
            assert quote_snapshot == {
                "items": [{"asin": "B07TC2BK1X", "quantity": 5, "unit_price": Decimal("121.98")}],
                "subtotal": Decimal("609.90"),
            }
        # A sticky offer banner intercepts the ordinary native click. Centering
        # the buy-box button lets the same browser click and confirm one add.
        html = '''<style>body {margin:0;height:2500px}
            #offer {position:fixed;top:0;left:0;width:100%;height:120px;z-index:10;background:#ddd}
            #add-to-cart-button {position:absolute;top:1000px;left:50px;width:180px;height:40px}
            </style><div id="offer">Offer display feature message</div>
            <span id="nav-cart-count">0</span>
            <button id="add-to-cart-button" onclick="document.getElementById('nav-cart-count').textContent='5'">Add to cart</button>'''
        driver.get("data:text/html," + quote(html))
        button = driver.find_element(By.ID, "add-to-cart-button")
        driver.execute_script("arguments[0].scrollIntoView({block:'start'});", button)
        with pytest.raises(ElementClickInterceptedException):
            button.click()
        add_amazon_item_to_cart(driver, button, "Part", tmp_path)
        assert driver.find_element(By.ID, "nav-cart-count").text == "5"

        # A persistent overlay must lead to manual recovery, not a JS click
        # through the overlay. Verify that the diagnostic is saved and the same
        # browser is available when the purchaser handles the offer prompt.
        driver.get("data:text/html," + quote(html))
        driver.execute_script("document.getElementById('offer').style.height='100vh';")
        button = driver.find_element(By.ID, "add-to-cart-button")

        def finish_manually(prompt):
            assert "Finish adding" in prompt
            assert driver.find_element(By.ID, "nav-cart-count").text == "0"
            assert (tmp_path / "cart_diagnostics" / "add_to_cart.png").exists()
            driver.execute_script("document.getElementById('offer').remove();")
            button.click()
            return ""

        monkeypatch.setattr("builtins.input", finish_manually)
        add_amazon_item_to_cart(driver, button, "Part", tmp_path)
        assert driver.find_element(By.ID, "nav-cart-count").text == "5"

        # The coverage dialog is rendered after Add to Cart, with Amazon-style
        # input/aria-labelledby markup. Declining is the only action that commits
        # the merchandise, and must never add the optional protection item.
        html = '''<span id="nav-cart-count">0</span>
            <button id="add-to-cart-button" onclick="document.getElementById('coverage').style.display='block'">Add to cart</button>
            <div id="coverage" class="a-popover" role="dialog" style="display:none">
            <h2>Add a protection plan?</h2><p>Optional coverage</p>
            <button id="add-protection" onclick="document.getElementById('nav-cart-count').textContent='6'">Add Protection</button>
            <span id="attachSiNoCoverage"><input type="button" aria-labelledby="attachSiNoCoverage-announce"
            onclick="document.getElementById('nav-cart-count').textContent='5';document.getElementById('coverage').style.display='none'">
            <span id="attachSiNoCoverage-announce">No Thanks</span></span></div>'''
        driver.get("data:text/html," + quote(html))
        button = driver.find_element(By.ID, "add-to-cart-button")
        add_amazon_item_to_cart(driver, button, "Part", tmp_path)
        assert driver.find_element(By.ID, "nav-cart-count").text == "5"
        assert not driver.find_element(By.ID, "coverage").is_displayed()
        # DigiKey's quantity change and native add are checked against a visible
        # cart counter. This local fixture does not pass a live bot challenge.
        from mrg_finance import digikey_cart
        html = '''<a href="https://www.digikey.com/ordering/shoppingcart" id="count">Your Item(s) 0</a>
            <div id="buybox"><input type="text" inputmode="numeric" value="1" max="150" id="quantity">
            <button id="add" onclick="this.dataset.clicks=String(Number(this.dataset.clicks||0)+1);
            document.getElementById('count').textContent='Your Item(s) '+document.getElementById('quantity').value;
            document.getElementById('status').textContent='Added to your cart'">Add to Cart</button>
            <div role="status" id="status"></div></div>
            <div class="related-products"><input type="number" value="1"><input type="number" value="1"></div>'''
        monkeypatch.setattr(digikey_cart, "navigate_for_evidence", lambda d, u: d.get("data:text/html," + quote(html)))
        monkeypatch.setattr(digikey_cart.price_scraper, "dismiss_popups_and_interstitials", lambda d: None)
        digikey_cart.add_digikey_item(driver, {"item_name": "Raspberry Pi", "quantity": 5,
            "link": "https://www.digikey.com/en/products/detail/raspberry-pi/SC0194-9/10258781"}, tmp_path)
        assert driver.find_element(By.ID, "quantity").get_attribute("value") == "5"
        assert driver.find_element(By.ID, "add").get_attribute("data-clicks") == "1"
        assert driver.find_element(By.ID, "count").text == "Your Item(s) 5"
        # Native quantity selection must be committed and blurred before the
        # buy-box click. A value that stays at 1 would add the wrong quantity.
        from mrg_finance import price_scraper
        html = '''<span id="nav-cart-count">0</span>
            <select id="quantity"><option value="1">1</option><option value="5">5</option></select>
            <button id="add-to-cart-button" onclick="document.getElementById('nav-cart-count').textContent=document.getElementById('quantity').value">Add to Cart</button>'''
        driver.get("data:text/html," + quote(html))
        assert price_scraper.check_and_set_amazon_quantity(driver, 5, "Part") == (5, False, "")
        assert driver.switch_to.active_element.get_attribute("id") != "quantity"
        add_amazon_item_to_cart(driver, driver.find_element(By.ID, "add-to-cart-button"), "Part", tmp_path)
        assert driver.find_element(By.ID, "nav-cart-count").text == "5"
    finally:
        driver.quit()
