import os
from pathlib import Path
from unittest.mock import MagicMock
from urllib.parse import quote

import openpyxl
import pytest
from selenium.webdriver.common.by import By

from mrg_finance import engage_fields, engage_bill_lookup, order_excel_builder
from mrg_finance.vendor_payee import lookup_vendor_payee


def control(id_, value=""):
    field = MagicMock()
    attributes = {"id": id_, "value": value}
    field.get_attribute.side_effect = attributes.get
    field.clear.side_effect = lambda: attributes.update(value="")
    field.send_keys.side_effect = lambda text: attributes.update(value=attributes["value"] + str(text))
    return field


def native_form():
    driver = MagicMock()
    fields = {key: control(key) for key in ("RequestedAmount", "refs", "budget-answer", "bill-answer",
              "PayeeFirstName", "PayeeLastName", "PayeeStreet", "PayeeStreet2", "PayeeCity", "PayeeState", "PayeeZipCode", "email")}
    groups = [{"kind": "budget", "selected": True, "field": fields["budget-answer"]},
              {"kind": "bill", "selected": False, "field": fields["bill-answer"]}]
    driver.find_elements.side_effect = lambda by, key: [fields[key]] if key in fields else []
    def script(js, labels=None, excluded=None):
        if labels == ["__sga_funding_groups__"]:
            return groups
        if labels and labels[0] == "requested amount":
            return [fields["RequestedAmount"]]
        if labels and labels[0].startswith("What is"):
            return [fields["refs"]]
        if labels and labels[0] == "payee email":
            return [fields["email"]]
        return []
    driver.execute_script.side_effect = script
    return driver, fields, groups


def fill_budget(driver):
    return engage_fields.fill_purchase_fields(driver, amount=600, cart_total=600,
        bill_refs="Budget 344042, B03, Line 1", sga_lines="$600.00, Budget 344042, B03, Line 1",
        funding_lines={"budget": "$600.00, Budget 344042, B03, Line 1"},
        payee=lookup_vendor_payee("DigiKey", ""))


def test_native_payee_and_selected_budget_answer_are_filled_without_touching_bill():
    driver, fields, _ = native_form()
    fill_budget(driver)
    assert fields["budget-answer"].get_attribute("value") == "$600.00, Budget 344042, B03, Line 1"
    fields["bill-answer"].clear.assert_not_called()
    fields["bill-answer"].send_keys.assert_not_called()
    for key, value in {"PayeeFirstName": "Digi-Key", "PayeeLastName": "Electronics",
                       "PayeeStreet": "701 Brooks Avenue South", "PayeeCity": "Thief River Falls",
                       "PayeeState": "MN", "PayeeZipCode": "56701", "email": "orders@digikey.com"}.items():
        assert fields[key].get_attribute("value") == value
    assert fields["refs"].get_attribute("value").startswith("Budget 344042")


def test_wrong_funding_checkbox_pauses_for_correction_in_same_form(monkeypatch, tmp_path):
    driver, fields, groups = native_form()
    groups[0]["selected"], groups[1]["selected"] = False, True
    monkeypatch.setattr(engage_fields, "save_page_screenshot", lambda *a: "diagnostic.png")
    def correct(prompt):
        driver.quit.assert_not_called()
        assert "retry" in prompt
        fields["budget-answer"].send_keys.assert_not_called()
        fields["bill-answer"].send_keys.assert_not_called()
        groups[0]["selected"], groups[1]["selected"] = True, False
        return ""
    monkeypatch.setattr("builtins.input", correct)
    engage_fields.run_engage_step(driver, tmp_path, "Form filling", lambda: fill_budget(driver))
    driver.get.assert_not_called()
    driver.quit.assert_not_called()
    assert fields["budget-answer"].get_attribute("value").startswith("$600.00")


def test_missing_payee_waits_for_user_and_retries_without_navigating(monkeypatch, tmp_path):
    driver, fields, groups = native_form()
    original = driver.find_elements.side_effect
    state = {"ready": False}
    driver.find_elements.side_effect = lambda by, key: [] if key.startswith("Payee") and not state["ready"] else original(by, key)
    monkeypatch.setattr(engage_fields, "save_page_screenshot", lambda *a: "diagnostic.png")
    def reveal(prompt):
        driver.quit.assert_not_called()
        state["ready"] = True
        return ""
    monkeypatch.setattr("builtins.input", reveal)
    engage_fields.run_engage_step(driver, tmp_path, "Form filling", lambda: fill_budget(driver))
    assert fields["PayeeFirstName"].get_attribute("value") == "Digi-Key"
    driver.get.assert_not_called()


def test_known_payee_email_cannot_be_silently_skipped(monkeypatch, tmp_path):
    driver, fields, _ = native_form()
    original = driver.execute_script.side_effect
    state = {"email_ready": False}
    driver.execute_script.side_effect = lambda js, labels=None, excluded=None: (
        [] if labels and labels[0] == "payee email" and not state["email_ready"] else original(js, labels, excluded))
    monkeypatch.setattr(engage_fields, "save_page_screenshot", lambda *a: "diagnostic.png")
    def ready(prompt):
        driver.quit.assert_not_called()
        state["email_ready"] = True
        return ""
    monkeypatch.setattr("builtins.input", ready)
    engage_fields.run_engage_step(driver, tmp_path, "Payee filling", lambda: fill_budget(driver))
    assert fields["email"].get_attribute("value") == "orders@digikey.com"
    driver.get.assert_not_called()


def test_recovery_closes_only_after_explicit_cancel(monkeypatch, tmp_path):
    driver = MagicMock()
    monkeypatch.setattr(engage_fields, "save_page_screenshot", lambda *a: "diagnostic.png")
    monkeypatch.setattr("builtins.input", lambda p: "cancel")
    action = MagicMock(side_effect=engage_fields.FormFieldError("Missing field"))
    with pytest.raises(SystemExit, match="cancelled"):
        engage_fields.run_engage_step(driver, tmp_path, "Form filling", action)
    driver.quit.assert_not_called()  # Caller closes on explicit cancellation only.


def test_lookup_classifies_live_annual_budget_title(monkeypatch):
    body = "Request: SGA FY27 Budget Request\nBudget Section:\nB03 - General Inventoried Goods\n1. Rapsberry Pi 4"
    monkeypatch.setattr(engage_bill_lookup, "open_budget_view", lambda d: body)
    result = engage_bill_lookup.lookup_bill_item_locations(MagicMock(), "344042", ["Rapsberry Pi 4"])
    assert result["Rapsberry Pi 4"]["funding_kind"] == "budget"


def test_comparison_identifies_budget_funding(tmp_path):
    rows = [{"bill_no": "344042", "item_name": "Part", "quantity": 6, "cost": 63.35,
             "quoted_unit_cost": 100, "quoted_total": 600, "funding_kind": "budget",
             "resolved_line_id": 1, "resolved_section": "B03"}]
    path, _ = order_excel_builder.generate_order_budget_vs_quoted_excel("budget-order", rows, output_dir=str(tmp_path))
    workbook = openpyxl.load_workbook(path)
    assert workbook.active["B5"].value == "Budget 344042"
    assert "Budget 344042" in workbook.active["A2"].value
    workbook.close()


def test_changed_controls_are_relocated_and_multiline_input_never_sends_enter():
    from selenium.common.exceptions import StaleElementReferenceException
    driver = MagicMock()
    state = {"field": control("refs"), "replaced": False}
    old = state["field"]
    old.tag_name = "input"

    def replace_on_clear():
        state["field"] = control("refs")
        state["field"].tag_name = "input"
        state["replaced"] = True
        old.get_attribute.side_effect = StaleElementReferenceException()

    old.clear.side_effect = replace_on_clear
    result = engage_fields.fill_and_verify(driver, lambda: state["field"], "Bill 1, Line 1\nBill 2, Line 3")
    assert state["replaced"]
    assert result["value"] == "Bill 1, Line 1; Bill 2, Line 3"
    old.send_keys.assert_not_called()
    state["field"].send_keys.assert_called_once_with("Bill 1, Line 1; Bill 2, Line 3")
    # A same-page retry should not erase successfully entered answers.
    engage_fields.fill_and_verify(driver, lambda: state["field"], result["value"])
    old.clear.assert_called_once()
    state["field"].clear.assert_not_called()


def test_report_shows_spreadsheet_section_before_lookup_without_inventing_line(tmp_path):
    rows = [{"bill_no": "344042", "item_name": "Part", "quantity": 2, "cost": 10,
             "budget_section": "B06 - Non-Inventoried Items", "bill_item_id": "999"}]
    path, _ = order_excel_builder.generate_order_budget_vs_quoted_excel("order", rows, output_dir=str(tmp_path))
    wb = openpyxl.load_workbook(path)
    assert wb.active["D5"].value == "B06 - Non-Inventoried Items"
    assert wb.active["A5"].value == "Unverified line"
    wb.close()


def test_standalone_report_passes_each_linked_bill_section_and_honors_explicit_path(monkeypatch, tmp_path):
    path = tmp_path / "custom-master.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bills"
    ws.append(["Bill Item ID", "Bill No.", "Bill Title", "Item Name", "Budget Section", "Cost", "Link"])
    ws.append([1, 344042, "FY27 Budget", "Part A", "B06 - Non-Inventoried Items", 10, ""])
    ws.append([2, 344042, "FY27 Budget", "Part B", "B03 - General Inventoried Goods", 20, ""])
    ws = wb.create_sheet("Ordering")
    ws.append(["Order ID", "Bill Item ID", "Quantity"])
    ws.append(["section-test", 1, 2])
    ws.append(["section-test", 2, 1])
    wb.save(path)
    wb.close()
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr("sys.argv", ["report", "--order", "section-test", "--excel-path", str(path), "--skip-scrape"])
    order_excel_builder.main()
    report = openpyxl.load_workbook(tmp_path / "screenshots/section-test/Budget_vs_Quoted_Detail_section-test.xlsx")
    assert report.active["D5"].value == "B06 - Non-Inventoried Items"
    assert report.active["D6"].value == "B03 - General Inventoried Goods"
    assert report.active["B5"].value == "Budget 344042"
    assert report.active["A5"].value == "Unverified line"
    report.close()


@pytest.mark.skipif(os.environ.get("MRG_RUN_BROWSER_TESTS") != "1", reason="Opt-in local Chrome DOM regression")
def test_real_dom_multiline_bill_reference_and_replaced_controls_do_not_submit(tmp_path):
    from selenium import webdriver
    from selenium.webdriver.chrome.options import Options
    options = Options()
    options.add_argument("--headless=new")
    driver = webdriver.Chrome(options=options)
    html = '''<form id="request">
      <label for="Subject">Subject</label><input id="Subject" value="Amazon purchase">
      <label for="Description">Description</label><textarea id="Description">https://share-a-cart.com/get/NOF2A</textarea>
      <label for="RequestedAmount">Requested Amount</label><input id="RequestedAmount" type="number">
      <div><label for="refs">What is the Budget/Bill # and Request Line #?</label><input id="refs"></div>
      <div><input id="bill-check" type="checkbox" checked><label for="bill-check">SGA Bill</label>
        <textarea id="answerTextBox-bill-check-free"></textarea></div>
      <section><h2>Payee Information</h2>'''
    for key in ("PayeeFirstName", "PayeeLastName", "PayeeStreet", "PayeeCity", "PayeeState", "PayeeZipCode"):
        html += f'<label for="{key}">{key}</label><input id="{key}">'
    html += '''</section><label for="email">Payee Email:</label><input id="email">
      <button type="submit">Submit</button></form><script>
      window.submissions=0; window.redraws=0;
      const form=document.querySelector('form');
      form.addEventListener('submit', e => { e.preventDefault(); window.submissions++; });
      form.addEventListener('change', e => {
        // Blur/clear replaces sibling inputs, like conditional Engage questions.
        for (const field of [...form.querySelectorAll('input,textarea')]) {
          if (field===e.target) continue;
          const replacement=field.cloneNode(true); replacement.value=field.value;
          field.replaceWith(replacement);
        }
        window.redraws++;
      });</script>'''
    refs = "Bill 376582, B06, Line 4\nBill 376582, B03, Line 2"
    try:
        driver.get("data:text/html," + quote(html))
        # Reproduce the previous Enter-key behavior in a local, non-submitting form.
        driver.find_element(By.ID, "refs").send_keys(refs)
        assert driver.execute_script("return window.submissions") > 0
        driver.execute_script("document.getElementById('refs').value='';window.submissions=0;")
        result = engage_fields.fill_purchase_fields(driver, amount=25, cart_total=25,
            bill_refs=refs, funding_lines={"bill": "$10, Line 4\n$15, Line 2"},
            payee=lookup_vendor_payee("DigiKey", ""))
        assert result == "25.00"
        assert driver.find_element(By.ID, "refs").get_attribute("value") == refs.replace("\n", "; ")
        assert driver.find_element(By.ID, "answerTextBox-bill-check-free").get_attribute("value") == "$10, Line 4\n$15, Line 2"
        assert driver.find_element(By.ID, "PayeeFirstName").get_attribute("value") == "Digi-Key"
        assert driver.find_element(By.ID, "PayeeCity").get_attribute("value") == "Thief River Falls"
        assert driver.execute_script("return window.submissions") == 0
        assert driver.execute_script("return window.redraws") > 0
        # Replace the amount control after filling, then read the fresh value.
        driver.execute_script("const field=document.getElementById('RequestedAmount');const fresh=field.cloneNode(true);fresh.value=field.value;field.replaceWith(fresh);")
        assert engage_fields.read_requested_amount(driver) == "25.00"
        assert driver.find_element(By.ID, "Subject").get_attribute("value") == "Amazon purchase"
        assert driver.find_element(By.ID, "Description").get_attribute("value").endswith("NOF2A")
    finally:
        driver.quit()


@pytest.mark.skipif(os.environ.get("MRG_RUN_BROWSER_TESTS") != "1", reason="Opt-in local Chrome DOM regression")
def test_real_dom_budget_checkboxes_and_native_payee_preserve_existing_form(tmp_path, monkeypatch):
    from selenium import webdriver
    from selenium.webdriver.chrome.options import Options
    options = Options()
    options.add_argument("--headless=new")
    driver = webdriver.Chrome(options=options)
    html = '''<form>
      <label for="Subject">Subject</label><input id="Subject" value="DigiKey purchase">
      <label for="Description">Description</label><textarea id="Description">https://share-a-cart.com/get/A0V3Q</textarea>
      <label for="RequestedAmount">Requested Amount</label><input id="RequestedAmount" type="number">
      <div><label for="refs">What is the Budget/Bill # and Request Line #?</label><input id="refs"></div>
      <div><input id="45927436" type="checkbox"><label for="45927436">SGA Budget<br>Include total $ reimbursement amount below</label>
        <label for="answerTextBox-45927436-free">Write-In Answer</label><textarea id="answerTextBox-45927436-free"></textarea></div>
      <div><input id="45927437" type="checkbox" checked><label for="45927437">SGA Bill<br>Include Bill # and total reimbursement amount below</label>
        <label for="answerTextBox-45927437-free">Write-In Answer</label><textarea id="answerTextBox-45927437-free"></textarea></div>
      <section><h2>Payee Information</h2>'''
    for key, label in (("PayeeFirstName", "First Name"), ("PayeeLastName", "Last Name"),
                       ("PayeeStreet", "Street"), ("PayeeCity", "City"), ("PayeeState", "State/Province"), ("PayeeZipCode", "ZIP/Postal Code")):
        html += f'<label for="{key}">{label}</label><input id="{key}">'
    html += '</section><label for="email">Payee&nbsp;Email:</label><input id="email"></form>'
    try:
        driver.get("data:text/html," + quote(html))
        def correct(prompt):
            assert driver.find_element(By.ID, "Subject").get_attribute("value") == "DigiKey purchase"
            assert driver.find_element(By.ID, "Description").get_attribute("value").endswith("A0V3Q")
            driver.find_element(By.ID, "45927437").click()
            driver.find_element(By.ID, "45927436").click()
            return ""
        monkeypatch.setattr("builtins.input", correct)
        engage_fields.run_engage_step(driver, tmp_path, "Form filling", lambda: fill_budget(driver))
        assert driver.find_element(By.ID, "answerTextBox-45927436-free").get_attribute("value").startswith("$600.00, Budget")
        assert driver.find_element(By.ID, "answerTextBox-45927437-free").get_attribute("value") == ""
        assert driver.find_element(By.ID, "PayeeCity").get_attribute("value") == "Thief River Falls"
        assert driver.find_element(By.ID, "PayeeFirstName").get_attribute("value") == "Digi-Key"
        assert driver.find_element(By.ID, "email").get_attribute("value") == "orders@digikey.com"
        assert driver.find_element(By.ID, "Subject").get_attribute("value") == "DigiKey purchase"
    finally:
        driver.quit()
