"""Replacement products must persist and rebuild every part of the vendor quote."""

from decimal import Decimal
from unittest.mock import MagicMock

import openpyxl
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table
import pytest

from mrg_finance import price_scraper, purchase_cart, share_a_cart
from mrg_finance.screenshot_capture import BrowserChallenge
from mrg_finance.purchase_sources import (
    CartReplacement, OrderingLinkStore, edit_links_before_cart,
    offer_quote_replacements, order_vendor, purchase_source,
)
from mrg_finance.spreadsheet_utils import read_sheet_robust


AMAZON = "https://www.amazon.com/dp/B012345678"
DIGIKEY = "https://www.digikey.com/en/products/detail/example/part/123"


def item(row=3, **kwargs):
    return {"item_name": "Part", "cost": 10, "quantity": 2, "link": AMAZON,
            "bill_item_id": "36", "bill_no": "344042", "ordering_row": row, **kwargs}


def workbook(path, header_row=2):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ordering"
    if header_row == 2:
        ws.append(["TOTALS"])
    ws.append(["Order ID (YYMMDD_vendor_gburdell3)", "Bill Item ID", "Product Link", "Vendor",
               "Cost", "Quantity", "Total Cost", "Share-A-Cart Link"])
    row = header_row + 1
    ws.append(["old_amazon_order", "36", AMAZON, '=XLOOKUP(B3,Bills!A:A,Bills!D:D)', 10, 2,
               f"=E{row}*F{row}", "https://share-a-cart.com/get/STALE"])
    ws.append(["another-order", "37", AMAZON, "Amazon", 5, 1, f"=E{row + 1}*F{row + 1}", "keep-cart"])
    ws.cell(row, 3).font = Font(color="0000FF", underline="single")
    ws.cell(row, 3).hyperlink = AMAZON
    ws.cell(row, 8).hyperlink = ws.cell(row, 8).value
    ws.add_table(Table(displayName="OrderT", ref=f"A{header_row}:H{row + 1}"))
    bills = wb.create_sheet("Bills")
    bills.append(["Bill Item ID", "Cost", "Link", "Vendor"])
    bills.append(["36", 10, AMAZON, "Amazon"])
    wb.save(path)
    wb.close()
    return row


@pytest.mark.parametrize("header_row", [1, 2])
def test_substitution_saves_only_ordering_inputs_and_preserves_bill_and_table(tmp_path, monkeypatch, header_row):
    path = tmp_path / "synced.xlsx"
    row = workbook(path, header_row)
    df = read_sheet_robust(str(path), ["Ordering"])
    assert df.attrs["header_row"] == header_row
    items = [item(row, quoted_unit_cost=12.5, quoted_total=25, asin="B012345678", resolved_url=AMAZON)]
    store = OrderingLinkStore(path, "old_amazon_order", items, header_row)
    answers = iter([DIGIKEY])
    monkeypatch.setattr("builtins.input", lambda p: next(answers))
    with pytest.raises(CartReplacement) as error:
        offer_quote_replacements(items, "Amazon", store)
    assert error.value.vendor == "DigiKey"
    assert items[0]["link"] == DIGIKEY
    assert items[0]["cost"] == 10 and items[0]["bill_no"] == "344042"
    assert not {"asin", "resolved_url", "quoted_unit_cost", "quoted_total"}.intersection(items[0])
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb["Ordering"]
    assert ws.cell(row, 3).value == DIGIKEY
    assert ws.cell(row, 3).hyperlink.target == DIGIKEY
    assert ws.cell(row, 4).value == "DigiKey"
    assert ws.cell(row, 7).value == f"=E{row}*F{row}"
    assert ws.cell(row, 8).value is None
    assert ws.cell(row, 8).hyperlink is None
    assert ws.cell(row, 3).font.color.rgb == "000000FF"
    assert ws.tables["OrderT"].ref == f"A{header_row}:H{row + 1}"
    assert ws.cell(row + 1, 8).value == "keep-cart"
    assert list(wb["Bills"].values)[1] == ("36", 10, AMAZON, "Amazon")
    wb.close()
    assert purchase_source({"Product Link": DIGIKEY, "Vendor": "Amazon"}, {"Link": AMAZON}) == (DIGIKEY, "DigiKey")


def test_can_switch_vendor_before_opening_amazon_and_save_again(tmp_path, monkeypatch):
    path = tmp_path / "synced.xlsx"
    workbook(path)
    items = [item()]
    store = OrderingLinkStore(path, "old_amazon_order", items, 2)
    answers = iter(["y", "1", "not a URL", "1", DIGIKEY, ""])
    monkeypatch.setattr("builtins.input", lambda p: next(answers))
    assert edit_links_before_cart(items, "Amazon", store) == "DigiKey"
    items[0]["quoted_unit_cost"] = 15
    answers = iter([AMAZON])
    with pytest.raises(CartReplacement) as error:
        offer_quote_replacements(items, "DigiKey", store)
    assert error.value.vendor == "Amazon"
    wb = openpyxl.load_workbook(path)
    assert wb["Ordering"].cell(3, 3).value == AMAZON
    assert wb["Ordering"].cell(3, 4).value == "Amazon"
    wb.close()


@pytest.mark.parametrize("column,value", [(2, "37"), (3, AMAZON + "?new-choice=1"), (6, 5)])
def test_save_refuses_concurrent_ordering_changes(tmp_path, column, value):
    path = tmp_path / "synced.xlsx"
    workbook(path)
    store = OrderingLinkStore(path, "old_amazon_order", [item()], 2)
    wb = openpyxl.load_workbook(path)
    wb["Ordering"].cell(3, column).value = value
    wb.save(path)
    wb.close()
    changed = path.read_bytes()
    with pytest.raises(ValueError, match="edited or moved"):
        store({0: DIGIKEY}, "DigiKey")
    assert path.read_bytes() == changed


def test_save_refuses_moved_columns(tmp_path):
    path = tmp_path / "synced.xlsx"
    workbook(path)
    store = OrderingLinkStore(path, "old_amazon_order", [item()], 2)
    wb = openpyxl.load_workbook(path)
    wb["Ordering"].cell(2, 3).value = "Notes"
    wb.save(path)
    wb.close()
    with pytest.raises(ValueError, match="columns changed"):
        store({0: DIGIKEY}, "DigiKey")


def test_atomic_save_preserves_a_concurrent_workbook_edit(monkeypatch, tmp_path):
    path = tmp_path / "synced.xlsx"
    workbook(path)
    store = OrderingLinkStore(path, "old_amazon_order", [item()], 2)
    save = openpyxl.workbook.workbook.Workbook.save
    concurrent = []

    def write_concurrently(wb, destination):
        save(wb, destination)
        latest = openpyxl.load_workbook(path)
        latest["Bills"].cell(3, 1).value = f"Another user's new row {len(concurrent)}"
        save(latest, path)
        latest.close()
        concurrent.append(path.read_bytes())

    monkeypatch.setattr(openpyxl.workbook.workbook.Workbook, "save", write_concurrently)
    with pytest.raises(ValueError, match="changed while saving"):
        store({0: DIGIKEY}, "DigiKey")
    assert len(concurrent) == 3
    assert path.read_bytes() == concurrent[-1]
    assert list(tmp_path.glob("*.xlsx")) == [path]


def test_replacement_save_reloads_and_preserves_a_one_time_sync_change(monkeypatch, tmp_path):
    path = tmp_path / "synced.xlsx"
    workbook(path)
    store = OrderingLinkStore(path, "old_amazon_order", [item()], 2)
    save = openpyxl.workbook.workbook.Workbook.save
    writes = []

    def sync_once(wb, destination):
        save(wb, destination)
        writes.append(destination)
        if len(writes) == 1:
            latest = openpyxl.load_workbook(path)
            latest["Bills"].cell(3, 1).value = "Preserve this synced row"
            save(latest, path)
            latest.close()

    monkeypatch.setattr(openpyxl.workbook.workbook.Workbook, "save", sync_once)
    store({0: DIGIKEY}, "DigiKey")
    assert len(writes) == 2
    latest = openpyxl.load_workbook(path)
    assert latest["Ordering"].cell(3, 3).value == DIGIKEY
    assert latest["Ordering"].cell(3, 4).value == "DigiKey"
    assert latest["Ordering"].cell(3, 8).value is None
    assert latest["Bills"].cell(3, 1).value == "Preserve this synced row"
    latest.close()
    assert list(tmp_path.glob("*.xlsx")) == [path]


def test_failed_save_does_not_mutate_purchase_links_or_quote(monkeypatch):
    items = [item(quoted_unit_cost=12.5, quoted_total=25)]
    store = MagicMock(side_effect=OSError("workbook locked"))
    monkeypatch.setattr("builtins.input", lambda p: DIGIKEY)
    with pytest.raises(OSError, match="locked"):
        offer_quote_replacements(items, "Amazon", store)
    assert items[0]["link"] == AMAZON
    assert items[0]["quoted_total"] == 25


def test_under_allocation_does_not_prompt_and_mixed_vendors_are_rejected(monkeypatch):
    items = [item(quoted_unit_cost=9)]
    prompt = MagicMock()
    monkeypatch.setattr("builtins.input", prompt)
    persist = MagicMock()
    offer_quote_replacements(items, "Amazon", persist)
    prompt.assert_not_called()
    persist.assert_not_called()
    with pytest.raises(ValueError, match="one vendor"):
        order_vendor(items + [item(link=DIGIKEY)], "Amazon")


def test_unreadable_product_page_can_defer_price_entry_to_the_actual_cart(monkeypatch):
    items = [item(quoted_unit_cost=None)]
    persist = MagicMock()
    monkeypatch.setattr("builtins.input", lambda prompt: "")
    offer_quote_replacements(items, "Amazon", persist)
    assert items[0]["quoted_unit_cost"] is None
    assert "quoted_total" not in items[0]
    assert "_reviewed_quote" not in items[0]
    persist.assert_not_called()


def test_cart_substitution_exits_before_fees_and_sharing_and_closes_old_browser(monkeypatch, tmp_path):
    monkeypatch.chdir(tmp_path)
    driver = MagicMock()
    monkeypatch.setattr(purchase_cart.webdriver, "Chrome", lambda **kw: driver)
    monkeypatch.setattr(purchase_cart, "Service", MagicMock())
    monkeypatch.setattr(purchase_cart, "navigate_for_evidence", lambda *a: None)
    monkeypatch.setattr(purchase_cart, "capture_evidence", lambda d, p, **kw: str(p))
    monkeypatch.setattr(price_scraper, "scrape_price_from_driver", lambda d: "$12.50")
    monkeypatch.setattr(purchase_cart, "read_amazon_cart", lambda d: {
        "items": [{"asin": "B012345678", "quantity": 2, "unit_price": Decimal("12.50")}],
        "subtotal": Decimal("25.00"),
    })
    monkeypatch.setattr(share_a_cart, "find_share_a_cart_extension", lambda: None)
    share = MagicMock()
    monkeypatch.setattr(share_a_cart, "prompt_for_share_a_cart", share)
    answers = iter([DIGIKEY])
    monkeypatch.setattr("builtins.input", lambda p: next(answers))
    items, store = [item()], MagicMock()
    with pytest.raises(CartReplacement):
        purchase_cart.prepare_cart(items, "Amazon", "test", "personal",
            quote_review=lambda rows, vendor: offer_quote_replacements(rows, vendor, store))
    driver.quit.assert_called_once()
    share.assert_not_called()
    assert list(answers) == []
    assert items[0]["link"] == DIGIKEY


@pytest.mark.parametrize("blocked", [False, True])
def test_substitution_is_saved_before_quantity_selection_or_any_add_click(monkeypatch, tmp_path, blocked):
    monkeypatch.chdir(tmp_path)
    path = tmp_path / "synced.xlsx"
    workbook(path)
    items = [item()]
    store = OrderingLinkStore(path, "old_amazon_order", items, 2)
    driver = MagicMock()
    monkeypatch.setattr(purchase_cart.webdriver, "Chrome", lambda **kw: driver)
    monkeypatch.setattr(purchase_cart, "Service", MagicMock())
    monkeypatch.setattr(purchase_cart, "navigate_for_evidence", MagicMock())
    monkeypatch.setattr(price_scraper, "dismiss_popups_and_interstitials", lambda d: None)
    monkeypatch.setattr(price_scraper, "scrape_price_from_driver", lambda d: "$121.98")
    capture = MagicMock(side_effect=BrowserChallenge("Challenge is stuck") if blocked else None)
    monkeypatch.setattr(purchase_cart, "capture_evidence", capture)
    quantity = MagicMock(side_effect=AssertionError("Quantity was selected before replacement review"))
    add = MagicMock(side_effect=AssertionError("Cart mutated before replacement review"))
    monkeypatch.setattr(price_scraper, "check_and_set_amazon_quantity", quantity)
    monkeypatch.setattr(purchase_cart, "add_amazon_item_to_cart", add)
    monkeypatch.setattr(share_a_cart, "find_share_a_cart_extension", lambda: None)
    prompts = []
    def answer(prompt):
        prompts.append(prompt)
        driver.quit.assert_not_called()
        return DIGIKEY
    monkeypatch.setattr("builtins.input", answer)
    with pytest.raises(CartReplacement) as error:
        purchase_cart.prepare_cart(items, "Amazon", "old_amazon_order",
            quote_review=lambda rows, vendor: offer_quote_replacements(rows, vendor, store))
    assert error.value.vendor == "DigiKey"
    assert len(prompts) == 1
    assert "current unit price" in prompts[0] if blocked else "substitute product URL" in prompts[0]
    quantity.assert_not_called()
    add.assert_not_called()
    driver.quit.assert_called_once()
    wb = openpyxl.load_workbook(path)
    assert wb["Ordering"].cell(3, 3).value == DIGIKEY
    assert wb["Ordering"].cell(3, 4).value == "DigiKey"
    assert wb["Bills"].cell(2, 2).value == 10
    wb.close()


@pytest.mark.parametrize("final_price,prompt_count", [(12.5, 1), (13, 2)])
def test_kept_overprice_is_reviewed_again_only_if_the_cart_price_changes(monkeypatch, final_price, prompt_count):
    items = [item(quoted_unit_cost=12.5)]
    prompt = MagicMock(return_value="")
    monkeypatch.setattr("builtins.input", prompt)
    persist = MagicMock()
    offer_quote_replacements(items, "Amazon", persist)
    items[0]["quoted_unit_cost"] = final_price
    offer_quote_replacements(items, "Amazon", persist)
    assert prompt.call_count == prompt_count
    persist.assert_not_called()


def test_failed_amazon_quantity_selection_is_not_reported_as_success(monkeypatch):
    driver, select, controller = MagicMock(), MagicMock(), MagicMock()
    controller.options = [MagicMock(), MagicMock()]
    for option, value in zip(controller.options, ("1", "5")):
        option.get_attribute.return_value = value
    select.get_attribute.return_value = "1"  # Selection was ignored by the browser.
    driver.find_elements.side_effect = [[], [select], []]
    monkeypatch.setattr("selenium.webdriver.support.ui.Select", lambda e: controller)
    monkeypatch.setattr("time.sleep", lambda n: None)
    selected, failed, reason = price_scraper.check_and_set_amazon_quantity(driver, 5, "Part")
    assert failed and selected != 5
    assert "Could not verify" in reason
