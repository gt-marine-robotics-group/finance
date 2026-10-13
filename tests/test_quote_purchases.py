"""Flight quotes must reach comparison/Engage without a cart or invalid navigation."""

from decimal import Decimal
from pathlib import Path
from unittest.mock import MagicMock

import openpyxl
import pytest

from mrg_finance import automation, order_excel_builder, purchase_cart
from mrg_finance.purchase_sources import is_product_url
from mrg_finance.quote_attachment import prompt_quote_attachment
from mrg_finance.screenshot_capture import validate_evidence


@pytest.mark.parametrize("link", ["", "-", "N/A", "nan", None])
def test_linkless_flight_uses_one_reviewed_quote_and_never_opens_vendor_browser(monkeypatch, tmp_path, link):
    monkeypatch.chdir(tmp_path)
    source = tmp_path / "Airline quote.pdf"
    original = b"%PDF-1.4\nVendor itinerary and quote\n%%EOF"
    source.write_bytes(original)
    answers = iter([str(source), "", "$1,250.50"])
    monkeypatch.setattr("builtins.input", lambda _: next(answers))
    browser = MagicMock(side_effect=AssertionError("A flight quote must not open a cart browser"))
    monkeypatch.setattr(purchase_cart.webdriver, "Chrome", browser)
    share = MagicMock(side_effect=AssertionError("A flight quote must not request Share-A-Cart"))
    monkeypatch.setattr(purchase_cart.share_a_cart, "prompt_for_share_a_cart", share)
    items = [{"item_name": "International RobotX Flights", "quantity": 1,
              "cost": 1200, "bill_no": "344042", "link": link}]
    cart = purchase_cart.prepare_cart(items, "Turkish Airlines", "flight-test")
    assert cart["driver"] is None and cart["share_url"] is None
    assert cart["total"] == Decimal("1250.50")
    assert items[0]["cost"] == 1200  # Preserve approved allocation.
    assert items[0]["quoted_unit_cost"] == 1250.5
    assert Path(cart["screenshot"]).read_bytes() == source.read_bytes() == original
    assert validate_evidence(cart["screenshot"])
    assert list(answers) == []  # Single-item quote total is entered only once.
    report, _ = order_excel_builder.generate_order_budget_vs_quoted_excel(
        "flight-test", items, output_dir=str(cart["folder"]), preliminary=True, cart_quote=cart)
    wb = openpyxl.load_workbook(report)
    assert wb.active["F5"].value == 1200
    assert wb.active["I5"].value == 1250.5
    summary = dict(wb["Cart Reconciliation"].values)
    assert summary["Requested amount"] == 1250.5
    assert "Manually reviewed" in summary["Quote basis"]
    wb.close()
    monkeypatch.setattr("builtins.input", lambda _: "")
    purchase_cart.recheck_cart(cart, items)
    Path(cart["screenshot"]).write_bytes(b"changed evidence")
    with pytest.raises(ValueError, match="changed or is missing"):
        purchase_cart.recheck_cart(cart, items)
    browser.assert_not_called()
    share.assert_not_called()


def test_empty_amazon_link_fails_clearly_without_launching_chrome(monkeypatch, tmp_path):
    monkeypatch.chdir(tmp_path)
    browser = MagicMock()
    monkeypatch.setattr(purchase_cart.webdriver, "Chrome", browser)
    with pytest.raises(ValueError, match="complete HTTP"):
        purchase_cart.prepare_cart([{"item_name": "Part", "link": "-"}], "Amazon", "test")
    browser.assert_not_called()


def test_purchase_command_accepts_linkless_flight_and_saves_preview_before_login(monkeypatch, tmp_path):
    from mrg_finance import automation_purchase
    monkeypatch.chdir(tmp_path)
    wb = openpyxl.Workbook()
    bills = wb.active
    bills.title = "Bills"
    bills.append(["Bill Item ID", "Bill No.", "Bill Title", "Item Name", "Vendor", "Cost", "Quantity", "Link"])
    bills.append(["1", "344042", "Flight Bill", "International RobotX Flights", "Turkish Airlines", 1000, 1, None])
    ordering = wb.create_sheet("Ordering")
    ordering.append(["Order ID", "Bill Item ID", "Quantity", "Status", "Vendor", "Link", "Cost"])
    ordering.append(["flight-test", "1", 1, "Pending", "Turkish Airlines", None, 1000])
    path = tmp_path / "budget.xlsx"
    wb.save(path)
    wb.close()
    source = tmp_path / "itinerary.pdf"
    source.write_bytes(b"%PDF-1.4\nCurrent flight quote\n%%EOF")
    monkeypatch.setattr(automation_purchase, "XLSX_PATH", str(path))
    monkeypatch.setattr(automation_purchase.sys, "argv", ["automation_purchase", "--order", "flight-test", "--excel-path", str(path)])
    monkeypatch.setattr(automation_purchase, "USERNAME", "")
    monkeypatch.setattr(automation_purchase, "PASSWORD", "")
    browser = MagicMock(side_effect=AssertionError("User stops before opening Engage"))
    monkeypatch.setattr(automation_purchase.webdriver, "Chrome", browser)
    answers = iter(["y", str(source), "", "1250", "n"])
    monkeypatch.setattr("builtins.input", lambda _: next(answers))
    with pytest.raises(SystemExit) as error:
        automation_purchase.main()
    assert error.value.code == 0
    assert list(answers) == []
    preview = openpyxl.load_workbook(tmp_path / "screenshots/flight-test/Budget_vs_Quoted_Detail_flight-test.xlsx")
    assert preview.active["F5"].value == 1000
    assert preview.active["I5"].value == 1250
    assert dict(preview["Cart Reconciliation"].values)["Requested amount"] == 1250
    preview.close()
    browser.assert_not_called()


def test_bill_quote_pdf_is_checked_and_can_be_resolved_for_attachment(monkeypatch, tmp_path):
    monkeypatch.setattr(automation, "SCREENSHOT_DIR", str(tmp_path / "screenshots"))
    source = tmp_path / "flight quote.pdf"
    source.write_bytes(b"%PDF-1.4\nQuote\n%%EOF")
    folder = tmp_path / "screenshots" / "Flight Bill"
    folder.mkdir(parents=True)
    (folder / "Flight.png").write_bytes(b"old unverified screenshot")
    answers = iter([str(source), ""])
    monkeypatch.setattr("builtins.input", lambda _: next(answers))
    path = prompt_quote_attachment(folder, "Flight")
    assert automation._verified_screenshot("Flight", "Flight Bill", "") == path
    automation.require_bill_evidence([("Flight", "")], "Flight Bill")
    source.write_bytes(b"Original file later changed")
    assert validate_evidence(path)  # The evidence copy is independent of its source.


def test_quote_file_prompt_recovers_from_missing_and_wrong_type_and_can_cancel(monkeypatch, tmp_path):
    wrong = tmp_path / "quote.pdf"
    wrong.write_bytes(b"<html>Verify you are human</html>")
    answers = iter([str(tmp_path / "missing.pdf"), str(wrong), "cancel"])
    monkeypatch.setattr("builtins.input", lambda _: next(answers))
    with pytest.raises(ValueError, match="cancelled"):
        prompt_quote_attachment(tmp_path / "evidence")
    assert not list((tmp_path / "evidence").iterdir())


@pytest.mark.parametrize("url", ["javascript:alert(1)", "file:///tmp/quote.pdf", "www.example.com", "https://[", "https://user:secret@example.com"])
def test_non_web_inputs_are_never_sent_to_chrome(url):
    assert not is_product_url(url)
