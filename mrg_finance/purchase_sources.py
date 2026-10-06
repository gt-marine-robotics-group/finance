"""Choose replacement products and keep Ordering links aligned with the purchase."""

from copy import deepcopy
import hashlib
from io import BytesIO
import os
from pathlib import Path
import tempfile
from urllib.parse import urlparse

import openpyxl

from .purchase_validation import money
from .spreadsheet_utils import COLUMN_ALIASES, clean_id, find_sheet_name, get_col_val
from .vendor_payee import vendor_key


def vendor_for_link(url, fallback=""):
    host = (urlparse(url).hostname or "").lower()
    domains = {
        "Amazon": ("amazon.com", "amazon.co.uk", "amazon.ca", "amazon.de", "a.co", "amzn.to"),
        "DigiKey": ("digikey.com", "digikey.ca", "digikey.co.uk"),
        "McMaster-Carr": ("mcmaster.com",), "Mouser": ("mouser.com",),
        "Adafruit": ("adafruit.com",), "SparkFun": ("sparkfun.com",),
        "Pololu": ("pololu.com",), "Blue Robotics": ("bluerobotics.com",),
    }
    for vendor, roots in domains.items():
        if any(host == root or host.endswith("." + root) for root in roots):
            return vendor
    if fallback and vendor_key(fallback) in vendor_key(host):
        return fallback
    return host.removeprefix("www.") or fallback


def purchase_source(row, bill_row, fallback=""):
    link = get_col_val(row, "link") or get_col_val(bill_row, "link")
    declared = get_col_val(row, "vendor") or get_col_val(bill_row, "vendor") or fallback
    return link, vendor_for_link(link, declared)


def order_vendor(items, fallback):
    vendors = [vendor_for_link(i["link"], fallback) for i in items]
    if len({vendor_key(v) for v in vendors}) != 1:
        raise ValueError("Each Order ID must contain one vendor. Replace the other links too, or split the order in Ordering.")
    return vendors[0]


def _replacement_url(raw):
    parsed = urlparse(raw)
    if parsed.scheme != "https" or not parsed.hostname or parsed.username or parsed.password:
        raise ValueError("Paste a complete HTTPS product URL.")
    return raw


class CartReplacement(Exception):
    def __init__(self, vendor):
        self.vendor = vendor
        super().__init__(f"Rebuild and verify the replacement cart with {vendor}.")


def _edit_links(proposed):
    while True:
        print("\nPurchase product links:")
        for n, item in enumerate(proposed, 1):
            print(f"  {n}. {item['item_name']}\n     {item['link']}")
        answer = input("Item number to replace (Enter when done, or 'cancel'): ").strip()
        if answer.lower() in ("cancel", "q", "quit"):
            raise ValueError("Product replacement cancelled; no Engage request was prepared.")
        if not answer:
            try:
                order_vendor(proposed, "")
                return
            except ValueError as error:
                print(error)
                continue
        if not answer.isdigit() or not 1 <= int(answer) <= len(proposed):
            print("Enter one of the listed item numbers.")
            continue
        item = proposed[int(answer) - 1]
        raw = input(f"Replacement product URL for {item['item_name']} (Enter to keep): ").strip()
        if raw.lower() in ("cancel", "q", "quit"):
            raise ValueError("Product replacement cancelled; no Engage request was prepared.")
        if raw:
            try:
                item["link"] = _replacement_url(raw)
            except ValueError as error:
                print(error)


def _apply_links(items, proposed, vendor, persist):
    new_vendor = order_vendor(proposed, vendor)
    changes = {n: item["link"] for n, item in enumerate(proposed) if item["link"] != items[n]["link"]}
    if not changes:
        return None
    persist(changes, new_vendor)  # A failed save must not proceed with an unsaved substitution.
    for n, link in changes.items():
        items[n]["link"] = link
    for item in items:
        for stale in ("asin", "resolved_url", "quoted_unit_cost", "quoted_total", "_reviewed_quote"):
            item.pop(stale, None)
    print(f"Purchase vendor: {new_vendor}. Rebuilding the cart and quote from the saved product links.")
    return new_vendor


def edit_links_before_cart(items, vendor, persist):
    if input("Change product links before building the cart? [y/N]: ").strip().lower() not in ("y", "yes"):
        return order_vendor(items, vendor)
    proposed = deepcopy(items)
    _edit_links(proposed)
    return _apply_links(items, proposed, vendor, persist) or order_vendor(items, vendor)


def offer_quote_replacements(items, vendor, persist):
    proposed = deepcopy(items)
    for n, item in enumerate(items):
        approved = money(item["cost"])
        if item.get("quoted_unit_cost") is None:
            print(f"\nCould not read the current price for {item['item_name']}. Approved: ${approved}/unit.\n"
                  f"Product: {item['link']}\n"
                  f"Enter the price for ONE unit; this order needs {item['quantity']}. "
                  "Press Enter to try reading it from the cart instead, or paste a replacement URL now.")
            while True:
                raw = input("Enter current unit price, or paste a replacement product URL (Enter to try the cart, or 'cancel'): ").strip()
                if not raw:
                    break  # No invented quote; the cart must supply it before purchase preparation.
                if raw.lower() in ("cancel", "q", "quit"):
                    raise ValueError("Product replacement cancelled; no Engage request was prepared.")
                try:
                    if raw.startswith("https://"):
                        proposed[n]["link"] = _replacement_url(raw)
                        break
                    quoted = money(raw)
                    item["quoted_unit_cost"] = float(quoted)
                    item["quoted_total"] = float(quoted * item["quantity"])
                    break
                except ValueError:
                    print("Enter the price shown by the vendor, a complete HTTPS product URL, or 'cancel'.")
            if proposed[n]["link"] != item["link"] or item.get("quoted_unit_cost") is None:
                continue
        quoted = money(item["quoted_unit_cost"])
        if quoted <= approved:
            continue
        reviewed = (item["link"], str(quoted))
        if item.get("_reviewed_quote") == reviewed:
            continue
        print(f"\n{item['item_name']}: approved ${approved}/unit, current ${quoted}/unit "
              f"(+${money((quoted - approved) * item['quantity'])} for {item['quantity']} items).\n"
              "Press Enter to keep this product at the higher price, paste a different product URL to replace it, "
              "or type 'cancel' to stop. Keeping it does not place an order or submit Engage.")
        while True:
            raw = input("Enter substitute product URL (Enter to keep current, or 'cancel'): ").strip()
            if raw.lower() in ("cancel", "q", "quit"):
                raise ValueError("Product replacement cancelled; no Engage request was prepared.")
            if not raw:
                item["_reviewed_quote"] = reviewed
                break
            try:
                proposed[n]["link"] = _replacement_url(raw)
                break
            except ValueError as error:
                print(error)
    try:
        order_vendor(proposed, vendor)
    except ValueError as error:
        print(error)
        _edit_links(proposed)
    new_vendor = _apply_links(items, proposed, vendor, persist)
    if new_vendor:
        raise CartReplacement(new_vendor)


class OrderingLinkStore:
    """Patch selected Ordering cells and refuse to overwrite concurrent row edits."""

    def __init__(self, path, order_id, items, header_row):
        self.path = Path(path).resolve()
        self.order_id = order_id
        self.items = items
        self.header_row = header_row
        wb = openpyxl.load_workbook(self.path, data_only=False)
        try:
            self.sheet = find_sheet_name(wb, ["Ordering", "Orders", "OrderT"])
            if not self.sheet:
                raise ValueError("Ordering sheet not found; cannot save replacement links.")
            ws = wb[self.sheet]
            self.headers = tuple(cell.value for cell in ws[header_row])
            self.columns = {}
            for col, cell in enumerate(ws[header_row], 1):
                header = str(cell.value or "").lower().strip()
                for key in ("order_id", "bill_item_id", "item_name", "link", "vendor", "share_cart_link"):
                    if header in COLUMN_ALIASES[key] or (key == "order_id" and header.startswith("order id")):
                        self.columns[key] = col
            self.original = {}
            for item in items:
                row = item["ordering_row"]
                if clean_id(ws.cell(row, self.columns["order_id"]).value) != order_id:
                    raise ValueError("Ordering rows changed while loading. Rerun before replacing links.")
                self.original[row] = {col: ws.cell(row, col).value for col in range(1, len(self.headers) + 1)}
        finally:
            wb.close()

    def __call__(self, changes, vendor):
        if "link" not in self.columns:
            raise ValueError("Add a Link column to Ordering before saving replacement product URLs.")
        for attempt in range(3):
            if self._save_once(changes, vendor):
                return
            if attempt < 2:
                print("Workbook changed during the save. Reloading it and checking the order before retrying…")
        raise ValueError("Workbook changed while saving on three attempts. Close Excel, wait for OneDrive sync to finish, "
                         "then rerun. Replacement links were not saved.")

    def _save_once(self, changes, vendor):
        # Hash and parse the same snapshot, even if OneDrive replaces the path.
        contents = self.path.read_bytes()
        digest = hashlib.sha256(contents).digest()
        wb = openpyxl.load_workbook(BytesIO(contents), data_only=False)
        temp = None
        try:
            ws = wb[self.sheet]
            if tuple(cell.value for cell in ws[self.header_row]) != self.headers:
                raise ValueError("Ordering columns changed during purchase. Rerun to reload them before saving replacements.")
            for row, expected in self.original.items():
                if any(ws.cell(row, col).value != value for col, value in expected.items()):
                    raise ValueError("Ordering rows were edited or moved during purchase. Rerun to reload them before saving replacements.")
            for n, link in changes.items():
                cell = ws.cell(self.items[n]["ordering_row"], self.columns["link"])
                cell.value = link
                cell.hyperlink = link
            for item in self.items:
                row = item["ordering_row"]
                if "vendor" in self.columns:
                    ws.cell(row, self.columns["vendor"]).value = vendor
                if "share_cart_link" in self.columns:
                    cell = ws.cell(row, self.columns["share_cart_link"])
                    cell.value = None
                    cell.hyperlink = None
            with tempfile.NamedTemporaryFile(dir=self.path.parent, suffix=".xlsx", delete=False) as file:
                temp = Path(file.name)
            wb.save(temp)
            if hashlib.sha256(self.path.read_bytes()).digest() != digest:
                return False  # Discard this stale copy; reload and revalidate all selected row values.
            os.replace(temp, self.path)
            for row in self.original:
                self.original[row] = {col: ws.cell(row, col).value for col in self.original[row]}
            print(f"Saved replacement URL(s) to Ordering in {self.path}")
            print("OneDrive will sync this workbook when connected if it is in your synced folder.")
            return True
        finally:
            wb.close()
            if temp:
                temp.unlink(missing_ok=True)
